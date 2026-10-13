from __future__ import annotations

import argparse
from dataclasses import replace
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import Error as PlaywrightError

from ehrm.browser.captcha_policy import is_allowed_host_url
from ehrm.browser.login import LoginService
from ehrm.browser.manager import BrowserManager
from ehrm.core.auth_repository import AuthenticationRepository, medical_browser_profile_name
from ehrm.core.exceptions import EhrmError
from ehrm.core.runtime import application_runtime_root, resolve_runtime_path
from ehrm.core.settings import DEFAULT_SETTINGS_PATH, load_settings
from ehrm.entrypoints.medical_termination_e2e_cli import (
    _MEDICAL_LOGIN_URL,
    _message,
    _save_screenshot,
    _select_account,
    _validate_login_url,
    _government_accounts,
)
from ehrm.modules.medical_enrollment.excel_loader import MedicalPreSubmitExcelLoader
from ehrm.modules.medical_enrollment.models import MedicalStage
from ehrm.modules.medical_enrollment.page import (
    MedicalPreSubmitBlockedError,
    MedicalPreSubmitPage,
)
from ehrm.modules.medical_termination.page import (
    is_medical_portal_ready,
    is_medical_portal_url,
    verify_existing_medical_session,
)
from ehrm.web.settings_service import WebSettingsService


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="ehrm-medical-enrollment-e2e",
        description=(
            "用现有医保账号测试人员基础信息采集或参保登记；"
            "只查询和填写录制可确认的字段，不保存、不提交。"
        ),
    )
    parser.add_argument("--config", type=Path, default=DEFAULT_SETTINGS_PATH)
    parser.add_argument(
        "--stage", choices=[stage.value for stage in MedicalStage],
        help="collection=人员基础信息采集；enrollment=人员参保登记",
    )
    parser.add_argument("--input", type=Path, help="每次只允许一名人员的 .xlsx/.xlsm 文件")
    parser.add_argument("--account-id", type=int, help="医保账号 ID；不填使用设置页的默认账号")
    parser.add_argument("--list-accounts", action="store_true", help="列出账号，不启动浏览器")
    parser.add_argument("--login-url", default=_MEDICAL_LOGIN_URL)
    parser.add_argument(
        "--output", type=Path, default=Path("output/medical-enrollment-e2e"),
        help="结果 Excel 与截图目录（相对于 runtime）",
    )
    parser.add_argument("--check-input", action="store_true", help="只检查账号与 Excel")
    parser.add_argument("--keep-browser-open", action="store_true", help="填写后等待人工检查")
    parser.add_argument("--ignore-https-errors", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        runtime_root = application_runtime_root(args.config)
        settings = load_settings(args.config, data_root=runtime_root)
        repository = AuthenticationRepository(settings.auth_database_path)
        if args.list_accounts:
            for account in _government_accounts(repository):
                print(f"ID={account.id}  {account.display_name or account.account}"
                      f"{'（默认）' if account.is_default else ''}")
            return 0
        if args.stage is None:
            raise ValueError("请通过 --stage 选择 collection 或 enrollment")
        if args.input is None:
            raise ValueError("请通过 --input 指定医保测试 Excel")
        stage = MedicalStage(args.stage)
        input_path = args.input.expanduser().resolve()
        item = MedicalPreSubmitExcelLoader().load(input_path, stage)
        account = _select_account(repository, args.account_id)
        login_url = _validate_login_url(args.login_url)
        settings = WebSettingsService(settings, runtime_root).runtime_settings(
            rights_account=account,
        )
        # Medical and social-security browser state must never share a profile.
        profile_root = runtime_root / "web-browser-profiles" / medical_browser_profile_name(account)
        settings = replace(
            settings,
            browser=replace(
                settings.browser,
                user_data_dir=profile_root,
                storage_state_path=profile_root / "session-state.json",
            ),
        )
        output_dir = resolve_runtime_path(args.output, runtime_root)
    except (EhrmError, OSError, ValueError) as exc:
        print(f"医保测试准备失败：{_message(exc)}")
        return 2

    label = "人员基础信息采集" if stage is MedicalStage.COLLECTION else "人员参保登记"
    print(f"测试步骤：{label}；账号：{account.display_name or account.account}")
    print("本次仅填写录制可确认字段；不会保存、提交，也不会标记已办理")
    if args.check_input:
        print("账号和 Excel 检查通过；未启动浏览器")
        return 0

    outcome = "失败"
    details = ""
    page = None
    try:
        with BrowserManager(
            settings.browser,
            ignore_https_errors=args.ignore_https_errors,
            stealth_enabled=(
                settings.captcha.stealth_enabled
                and is_allowed_host_url(
                    settings.site.login_url, settings.captcha.allowed_hosts
                )
            ),
        ) as browser:
            page = browser.page
            try:
                print("阶段 1/2：检查医保会话，必要时通过政务网登录……")
                page = LoginService(page, settings, progress_callback=print).ensure_external_government_authenticated(
                    entry_url=login_url,
                    username=account.account,
                    password=account.password,
                    target_matches=is_medical_portal_url,
                    target_ready=is_medical_portal_ready,
                    existing_session_ready=verify_existing_medical_session,
                )
                print(f"阶段 2/2：{label}，只执行到保存或提交之前……")
                MedicalPreSubmitPage(
                    page, stage,
                    timeout_ms=settings.browser.action_timeout_ms,
                    pacing=settings.browser.pacing,
                    progress_callback=print,
                ).fill_item(item)
                outcome = "已填写待核对（未提交）"
                details = "只填写录制可确认字段；未保存、未提交，其他字段须人工核对"
            except MedicalPreSubmitBlockedError as exc:
                outcome = exc.outcome
                details = _message(exc)
                print(f"医保业务阻断：{details}")
            except (EhrmError, PlaywrightError, OSError, ValueError, EOFError) as exc:
                details = _message(exc)
                print(f"医保测试失败：{details}")
            finally:
                _save_screenshot(page, output_dir / f"medical-{stage.value}-e2e.png")
                if args.keep_browser_open:
                    try:
                        input("浏览器保持打开；请勿误点保存或提交。检查后按回车关闭：")
                    except EOFError:
                        print("标准输入已关闭，浏览器将正常退出")
    except (EhrmError, PlaywrightError, OSError, ValueError) as exc:
        details = _message(exc)
        print(f"医保测试失败：{details}")

    try:
        result = _write_result(input_path, output_dir, stage, outcome, details)
        print(f"结果 Excel：{result.resolve()}")
    except (OSError, ValueError) as exc:
        print(f"医保测试结果写入失败：{_message(exc)}")
        return 1
    return 0 if outcome == "已填写待核对（未提交）" else 1


def _write_result(
    source: Path, output_dir: Path, stage: MedicalStage, outcome: str, details: str,
) -> Path:
    workbook = load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    try:
        sheet = workbook.active
        column = sheet.max_column + 1
        for offset, heading in enumerate(("测试步骤", "处理结果", "提示/错误信息", "是否提交")):
            sheet.cell(1, column + offset, heading)
        for offset, value in enumerate((stage.value, outcome, details, "否")):
            sheet.cell(2, column + offset, value)
        sheet.column_dimensions[sheet.cell(1, column + 2).column_letter].width = 65
        output_dir.mkdir(parents=True, exist_ok=True)
        destination = output_dir / (
            f"medical-{stage.value}-{datetime.now():%Y%m%d-%H%M%S-%f}"
            f"{source.suffix.lower()}"
        )
        workbook.save(destination)
        return destination
    finally:
        workbook.close()


if __name__ == "__main__":
    raise SystemExit(main())
