from __future__ import annotations

import argparse
from dataclasses import replace
from datetime import datetime
from pathlib import Path
import re
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook
from playwright.sync_api import Error as PlaywrightError

from ehrm.browser.captcha_policy import is_allowed_host_url
from ehrm.browser.login import LoginService
from ehrm.browser.manager import BrowserManager
from ehrm.core.auth_repository import (
    AuthenticationRepository, SystemAccount, SystemType, medical_browser_profile_name,
)
from ehrm.core.exceptions import EhrmError
from ehrm.core.runtime import application_runtime_root, resolve_runtime_path
from ehrm.core.settings import DEFAULT_SETTINGS_PATH, load_settings
from ehrm.modules.medical_termination.excel_loader import MedicalTerminationExcelLoader
from ehrm.modules.medical_termination.page import (
    MEDICAL_LOGIN_URL,
    MedicalTerminationPage,
    is_medical_portal_url,
    is_medical_portal_ready,
    verify_existing_medical_session,
)
from ehrm.web.settings_service import WebSettingsService


_GOVERNMENT_SYSTEM = "医保政务服务平台账号"
_MEDICAL_LOGIN_URL = MEDICAL_LOGIN_URL


def _validate_login_url(value: str) -> str:
    parts = urlsplit(value)
    query = parse_qs(parts.query)
    if (
        parts.scheme != "https"
        or parts.hostname != "www.jszwfw.gov.cn"
        or parts.path != "/jsjis/front/login.do"
        or set(query) != {"uuid"}
        or len(query["uuid"]) != 1
        or not query["uuid"][0].strip()
        or parts.fragment
        or parts.username is not None
        or parts.password is not None
    ):
        raise ValueError("医保政务网入口必须是带 uuid 的官方 HTTPS 登录地址")
    return value


def _government_accounts(repository: AuthenticationRepository) -> list[SystemAccount]:
    return [
        account for account in repository.list_accounts(SystemType.MEDICAL)
        if account.profile.get("login_method") == "government_sso"
    ]


def _select_account(
    repository: AuthenticationRepository, account_id: int | None
) -> SystemAccount:
    accounts = _government_accounts(repository)
    account = (
        next((item for item in accounts if item.id == account_id), None)
        if account_id is not None
        else next((item for item in accounts if item.is_default), None)
    )
    if account is None and account_id is None and len(accounts) == 1:
        account = accounts[0]
    if account is None:
        raise ValueError(
            f"系统设置中尚未配置“{_GOVERNMENT_SYSTEM}”，未设置默认账号，或指定账号 ID 无效"
        )
    if not account.account or not account.password:
        raise ValueError("所选政务平台账号信息不完整")
    return account


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="ehrm-medical-termination-e2e",
        description=(
            "使用已保存的政务平台账号登录，从 Excel 读取一名人员，"
            "测试医保停保查询和填表；绝不点击提交。"
        ),
    )
    parser.add_argument("--config", type=Path, default=DEFAULT_SETTINGS_PATH)
    parser.add_argument("--input", type=Path, help="只接受一名人员的医保停保 Excel")
    parser.add_argument("--account-id", type=int, help="医保账号 ID；不填使用设置页的默认账号")
    parser.add_argument("--list-accounts", action="store_true", help="列出政务平台账号，不启动浏览器")
    parser.add_argument(
        "--login-url", default=_MEDICAL_LOGIN_URL,
        help="医保专用的江苏政务服务网登录入口（含 uuid）",
    )
    parser.add_argument(
        "--output", type=Path,
        default=Path("output/medical-termination-e2e"),
        help="结果 Excel 与截图的输出目录（相对于 runtime）",
    )
    parser.add_argument("--check-input", action="store_true", help="只检查账号与 Excel")
    parser.add_argument(
        "--keep-browser-open", action="store_true",
        help="结束后等待回车，便于人工核对页面",
    )
    parser.add_argument(
        "--ignore-https-errors", action="store_true",
        help="仅用于自签名证书测试环境",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        runtime_root = application_runtime_root(args.config)
        settings = load_settings(args.config, data_root=runtime_root)
        repository = AuthenticationRepository(settings.auth_database_path)
        if args.list_accounts:
            for candidate in _government_accounts(repository):
                print(f"ID={candidate.id}  {candidate.display_name or candidate.account}"
                      f"{'（默认）' if candidate.is_default else ''}")
            return 0
        if args.input is None:
            raise ValueError("请通过 --input 指定医保停保 Excel")
        login_url = _validate_login_url(args.login_url)
        input_path = args.input.expanduser().resolve()
        item = MedicalTerminationExcelLoader().load(input_path)
        account = _select_account(repository, args.account_id)
        settings = WebSettingsService(settings, runtime_root).runtime_settings(
            rights_account=account,
        )
        profile_root = runtime_root / "web-browser-profiles" / medical_browser_profile_name(account)
        settings = replace(
            settings,
            browser=replace(
                settings.browser,
                user_data_dir=profile_root,
                storage_state_path=profile_root / "session-state.json",
            ),
        )
        output_dir = resolve_runtime_path(args.output, runtime_root)
    except (EhrmError, OSError, ValueError) as exc:
        _failure("医保停保测试准备失败", exc)
        return 2

    print(f"测试账号：{account.display_name or account.account}（政务服务平台）")
    print(f"测试人员：第 {item.source_index} 行；仅查询和填表，不提交")
    if args.check_input:
        print("账号和 Excel 检查通过；本步骤未启动浏览器")
        return 0

    page = None
    outcome = "失败"
    details = ""
    try:
        with BrowserManager(
            settings.browser,
            ignore_https_errors=args.ignore_https_errors,
            stealth_enabled=(
                settings.captcha.stealth_enabled
                and is_allowed_host_url(
                    settings.site.login_url, settings.captcha.allowed_hosts
                )
            ),
        ) as browser:
            page = browser.page
            try:
                print("阶段 1/2：正在检查医保登录状态，必要时通过政务网登录……")
                login = LoginService(page, settings, progress_callback=print)
                page = login.ensure_external_government_authenticated(
                    entry_url=login_url,
                    username=account.account,
                    password=account.password,
                    target_matches=is_medical_portal_url,
                    target_ready=is_medical_portal_ready,
                    existing_session_ready=verify_existing_medical_session,
                )
                print("已进入南京医保单位大厅")

                print("阶段 2/2：正在查询医保人员并填写停保信息……")
                MedicalTerminationPage(
                    page,
                    timeout_ms=settings.browser.action_timeout_ms,
                    pacing=settings.browser.pacing,
                    progress_callback=print,
                ).fill_item(item)
                outcome = "录入成功（未提交）"
                print("医保停保填表测试完成；未点击提交")
            except (EhrmError, PlaywrightError, OSError, ValueError, EOFError) as exc:
                details = _message(exc)
                _failure("医保停保测试失败", exc)
            finally:
                _save_screenshot(page, output_dir / "medical-termination-e2e.png")
                if args.keep_browser_open:
                    input("浏览器保持打开，核对完毕后按回车关闭：")
    except (EhrmError, PlaywrightError, OSError, ValueError) as exc:
        details = _message(exc)
        _failure("医保停保测试失败", exc)

    try:
        result = _write_result(input_path, output_dir, outcome, details)
        print(f"结果 Excel：{result.resolve()}")
    except (OSError, ValueError) as exc:
        _failure("医保停保结果写入失败", exc)
        return 1
    return 0 if outcome == "录入成功（未提交）" else 1


def _write_result(source: Path, output_dir: Path, outcome: str, details: str) -> Path:
    workbook = load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    try:
        sheet = workbook.active
        result_column = sheet.max_column + 1
        sheet.cell(1, result_column, "处理结果")
        sheet.cell(1, result_column + 1, "错误信息")
        sheet.cell(2, result_column, outcome)
        sheet.cell(2, result_column + 1, details)
        sheet.column_dimensions[sheet.cell(1, result_column).column_letter].width = 20
        sheet.column_dimensions[sheet.cell(1, result_column + 1).column_letter].width = 60
        output_dir.mkdir(parents=True, exist_ok=True)
        destination = output_dir / (
            f"medical-termination-{datetime.now():%Y%m%d-%H%M%S-%f}"
            f"{source.suffix.lower()}"
        )
        workbook.save(destination)
        return destination
    finally:
        workbook.close()


def _save_screenshot(page: object | None, path: Path) -> None:
    if page is None:
        return
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        page.screenshot(path=str(path), full_page=True)
        print(f"页面截图：{path.resolve()}")
    except (AttributeError, OSError, PlaywrightError) as exc:
        print(f"页面截图保存失败：{exc}")


def _message(exc: Exception) -> str:
    message = str(exc)
    details = getattr(exc, "details", None)
    combined = f"{message}；{details}" if details and details != message else message
    return re.sub(
        r"(?i)(accessToken|refreshToken)=([^&#\s]+)",
        r"\1=[已隐藏]",
        combined,
    )


def _failure(label: str, exc: Exception) -> None:
    print(f"{label}：{_message(exc)}")


if __name__ == "__main__":
    raise SystemExit(main())
