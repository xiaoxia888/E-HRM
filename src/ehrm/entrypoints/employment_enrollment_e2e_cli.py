from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from playwright.sync_api import Error as PlaywrightError

from ehrm.browser.captcha_policy import is_allowed_host_url
from ehrm.browser.login import LoginService
from ehrm.browser.manager import BrowserManager
from ehrm.core.auth_repository import AuthenticationRepository, SystemType
from ehrm.core.exceptions import EhrmError, ExcelValidationError, QueryValidationError
from ehrm.core.logging import configure_logging
from ehrm.core.runtime import application_runtime_root, resolve_runtime_path
from ehrm.core.settings import DEFAULT_SETTINGS_PATH, load_settings
from ehrm.modules.employment_enrollment.models import (
    EmploymentEnrollmentItem,
    EmploymentEnrollmentPreparation,
    EmploymentEnrollmentResult,
)
from ehrm.modules.employment_enrollment.service import EmploymentEnrollmentService
from ehrm.modules.jshrss_hall import validate_nanjing_contract


_HEADERS = (
    "身份证号",
    "合同增加原因",
    "合同类别",
    "劳动合同开始日期",
    "劳动合同终止日期",
    "岗位工种",
    "参保身份",
    "月缴费工资",
)


def load_test_excel(path: Path) -> list[EmploymentEnrollmentItem]:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ExcelValidationError("参保测试文件仅支持 .xlsx 或 .xlsm")
    if not path.is_file():
        raise ExcelValidationError(f"参保测试文件不存在：{path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelValidationError("无法读取参保测试 Excel", details=str(exc)) from exc
    try:
        rows = workbook.active.iter_rows(values_only=True)
        raw_headers = next(rows, None)
        if raw_headers is None:
            raise ExcelValidationError("参保测试 Excel 为空")
        headers = [str(value).strip() if value is not None else "" for value in raw_headers]
        missing = [header for header in _HEADERS if header not in headers]
        if missing:
            raise ExcelValidationError("参保测试 Excel 缺少必要列：" + "、".join(missing))
        indexes = {header: headers.index(header) for header in _HEADERS}
        items: list[EmploymentEnrollmentItem] = []
        errors: list[str] = []
        for row_number, values in enumerate(rows, start=2):
            raw_identity = _value(values, indexes["身份证号"])
            if isinstance(raw_identity, (int, float)) and not isinstance(raw_identity, bool):
                errors.append(f"第 {row_number} 行身份证号必须设置为文本格式，避免 Excel 舍入")
                continue
            row = {header: _value(values, indexes[header]) for header in _HEADERS}
            if not any(value is not None and str(value).strip() for value in row.values()):
                continue
            try:
                items.append(
                    EmploymentEnrollmentItem(
                        identity_number=_text(row["身份证号"]),
                        contract_add_reason=_text(row["合同增加原因"]),
                        contract_type=_text(row["合同类别"]),
                        contract_start_date=row["劳动合同开始日期"],
                        contract_end_date=row["劳动合同终止日期"],
                        position_type=_text(row["岗位工种"]),
                        insured_identity=_text(row["参保身份"]) or "企业职工",
                        monthly_wage=row["月缴费工资"],
                        source_index=row_number,
                    ).normalized()
                )
            except QueryValidationError as exc:
                errors.append(str(exc))
        if errors:
            raise ExcelValidationError(
                "参保测试 Excel 数据校验失败", details="\n".join(errors)
            )
        if not items:
            raise ExcelValidationError("参保测试 Excel 中没有可执行的数据")
        return items
    finally:
        workbook.close()


def _value(values: tuple[object, ...], index: int) -> object | None:
    return values[index] if index < len(values) else None


def _text(value: object | None) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="ehrm-employment-enrollment-e2e",
        description="读取 Excel、复用已保存账号登录并填写参保表单；绝不点击确认提交。",
    )
    parser.add_argument("--config", type=Path, default=DEFAULT_SETTINGS_PATH)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("output/employment-enrollment-e2e.xlsx"),
    )
    parser.add_argument(
        "--diagnostic",
        type=Path,
        default=Path("output/employment-enrollment-e2e.png"),
    )
    parser.add_argument("--ignore-https-errors", action="store_true")
    parser.add_argument("--pause-after-fill", action="store_true")
    parser.add_argument("--check-input", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        runtime_root = application_runtime_root(args.config)
        settings = load_settings(args.config, data_root=runtime_root)
        validate_nanjing_contract(settings)
        input_path = args.input.expanduser().resolve()
        output_path = resolve_runtime_path(args.output, runtime_root)
        diagnostic_path = resolve_runtime_path(args.diagnostic, runtime_root)
        items = load_test_excel(input_path)
        account = AuthenticationRepository(settings.auth_database_path).get_default_account(
            SystemType.JSHRSS
        )
        if account is None or not account.secondary_account or not account.password:
            raise ValueError("SQLite 中没有完整的默认智慧人社单位账号，请先在系统设置中维护")
    except (EhrmError, OSError, ValueError) as exc:
        _print_failure("参保 E2E 准备失败", exc)
        return 2

    print(f"执行地区：南京（areaCode=320100，运行时强校验），共 {len(items)} 条")
    print("安全边界：只录入参保数据，不点击确认提交")
    if args.check_input:
        print("参保 E2E 配置、账号和 Excel 检查通过")
        return 0

    logger = configure_logging(runtime_root / "logs")
    try:
        with BrowserManager(
            settings.browser,
            ignore_https_errors=args.ignore_https_errors,
            stealth_enabled=(
                settings.captcha.stealth_enabled
                and is_allowed_host_url(settings.site.login_url, settings.captcha.allowed_hosts)
            ),
        ) as browser:
            page = browser.page
            try:
                print("阶段 1/2：复用现有智慧人社自动化登录")
                login = LoginService(page, settings, progress_callback=print)
                login.ensure_authenticated()
                page = login.page
                print("阶段 2/2：校验南京地区并录入参保信息")
                completed: list[EmploymentEnrollmentResult] = []

                def save_row(row: EmploymentEnrollmentResult) -> None:
                    completed.append(row)
                    _write_results(output_path, completed)

                result = EmploymentEnrollmentService(
                    settings,
                    logger,
                    progress_callback=print,
                    result_callback=save_row,
                ).prepare_with_page(page, items)
                _save_screenshot(page, diagnostic_path)
                print(
                    f"参保录入完成：总计 {result.total_count} 条，成功 {result.prepared_count} 条，"
                    f"存在问题 {result.total_count - result.prepared_count} 条，submitted=false"
                )
                print(f"结果 Excel：{output_path.resolve()}")
                if args.pause_after_fill:
                    input("页面停在提交前，检查完毕后按回车关闭浏览器：")
                return 0
            except (EhrmError, OSError, PlaywrightError, ValueError):
                _save_screenshot(page, diagnostic_path)
                raise
    except (EhrmError, OSError, PlaywrightError, ValueError) as exc:
        _print_failure("完整参保数据录入 E2E 失败", exc)
        return 1


def _write_results(path: Path, results: list[EmploymentEnrollmentResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "录入结果"
    sheet.append(["Excel行号", "身份证号", "结果", "问题编码", "问题信息", "是否提交"])
    for result in results:
        sheet.append(
            [
                result.item.source_index,
                result.item.identity_number,
                "已录入" if result.success else "存在问题",
                result.code,
                result.message,
                "否",
            ]
        )
    workbook.save(path)
    workbook.close()


def _save_screenshot(page: object, path: Path) -> None:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        page.screenshot(path=str(path), full_page=True)
        print(f"页面截图：{path.resolve()}")
    except (AttributeError, OSError, PlaywrightError) as exc:
        print(f"页面截图保存失败：{exc}")


def _print_failure(label: str, exc: Exception) -> None:
    message = str(exc)
    details = getattr(exc, "details", None)
    if details and details != message:
        message += f"；{details}"
    print(f"{label}：{message}")


if __name__ == "__main__":
    raise SystemExit(main())
