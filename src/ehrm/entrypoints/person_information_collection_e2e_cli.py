from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from playwright.sync_api import Error as PlaywrightError

from ehrm.browser.captcha_policy import is_allowed_host_url
from ehrm.browser.login import LoginService
from ehrm.browser.manager import BrowserManager
from ehrm.core.auth_repository import AuthenticationRepository, SystemType
from ehrm.core.exceptions import EhrmError, ExcelValidationError, QueryValidationError
from ehrm.core.logging import configure_logging
from ehrm.core.runtime import application_runtime_root, resolve_runtime_path
from ehrm.core.settings import DEFAULT_SETTINGS_PATH, load_settings
from ehrm.modules.person_information_collection.models import (
    PersonInformationItem,
    PersonInformationPreparation,
    PersonInformationResult,
)
from ehrm.modules.person_information_collection.page import validate_nanjing_contract
from ehrm.modules.person_information_collection.service import PersonInformationService


_HEADERS = (
    "身份证号", "姓名", "手机号", "民族", "户籍性质", "省", "市", "区县", "街道", "社区村"
)


def load_test_excel(path: Path) -> list[PersonInformationItem]:
    """Excel belongs only to this E2E adapter, never the application service."""
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ExcelValidationError("采集测试文件仅支持 .xlsx 或 .xlsm")
    if not path.is_file():
        raise ExcelValidationError(f"采集测试文件不存在：{path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelValidationError("无法读取采集测试 Excel", details=str(exc)) from exc
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise ExcelValidationError("采集测试 Excel 为空")
        header_names = [str(value).strip() if value is not None else "" for value in headers]
        missing = [name for name in _HEADERS if name not in header_names]
        if missing:
            raise ExcelValidationError("采集测试 Excel 缺少必要列：" + "、".join(missing))
        indexes = {name: header_names.index(name) for name in _HEADERS}
        items: list[PersonInformationItem] = []
        errors: list[str] = []
        for row_number, values in enumerate(rows, start=2):
            raw_identity = (
                values[indexes["身份证号"]]
                if indexes["身份证号"] < len(values)
                else None
            )
            if isinstance(raw_identity, (int, float)) and not isinstance(raw_identity, bool):
                errors.append(f"第 {row_number} 行身份证号必须设置为文本格式，避免 Excel 舍入")
                continue
            data = {
                name: _cell_text(values[indexes[name]] if indexes[name] < len(values) else None)
                for name in _HEADERS
            }
            if not any(data.values()):
                continue
            try:
                items.append(
                    PersonInformationItem(
                        identity_number=data["身份证号"],
                        name=data["姓名"],
                        mobile=data["手机号"],
                        nation=data["民族"],
                        household_type=data["户籍性质"],
                        address=(
                            data["省"], data["市"], data["区县"],
                            data["街道"], data["社区村"],
                        ),
                        source_index=row_number,
                    ).normalized()
                )
            except QueryValidationError as exc:
                errors.append(str(exc))
        if errors:
            raise ExcelValidationError("采集测试 Excel 数据校验失败", details="\n".join(errors))
        if not items:
            raise ExcelValidationError("采集测试 Excel 中没有可执行的数据")
        return items
    finally:
        workbook.close()


def _cell_text(value: object | None) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="ehrm-person-information-collection-e2e",
        description="从 Excel 读取测试数据、复用已保存账号登录并填写采集表单；绝不提交。",
    )
    parser.add_argument("--config", type=Path, default=DEFAULT_SETTINGS_PATH)
    parser.add_argument("--input", type=Path, required=True, help="包含十列采集数据的 .xlsx/.xlsm 文件")
    parser.add_argument(
        "--output", type=Path, default=Path("output/person-information-collection-e2e.xlsx"),
        help="逐行结果 Excel 路径",
    )
    parser.add_argument(
        "--diagnostic", type=Path, default=Path("output/person-information-collection-e2e.png"),
        help="成功或失败时保存的当前页面截图",
    )
    parser.add_argument("--ignore-https-errors", action="store_true")
    parser.add_argument("--pause-after-fill", action="store_true", help="录入后等待人工检查，仍不提交")
    parser.add_argument("--check-input", action="store_true", help="仅校验账号、配置与 Excel")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        runtime_root = application_runtime_root(args.config)
        settings = load_settings(args.config, data_root=runtime_root)
        validate_nanjing_contract(settings)
        input_path = args.input.expanduser().resolve()
        output_path = resolve_runtime_path(args.output, runtime_root)
        diagnostic_path = resolve_runtime_path(args.diagnostic, runtime_root)
        items = load_test_excel(input_path)
        account = AuthenticationRepository(settings.auth_database_path).get_default_account(
            SystemType.JSHRSS
        )
        if account is None or not account.secondary_account or not account.password:
            raise ValueError("SQLite 中没有完整的默认智慧人社单位账号，请先在系统设置中维护")
    except (EhrmError, OSError, ValueError) as exc:
        _print_failure("采集 E2E 准备失败", exc)
        return 2

    print(f"执行地区：南京（areaCode=320100，运行时强校验），共 {len(items)} 条")
    print("登录凭据：读取应用 SQLite 中的默认智慧人社账号")
    print("安全边界：每人独立表单标签页，只录入、不点击业务提交")
    if args.check_input:
        print("采集 E2E 配置、账号和 Excel 检查通过")
        return 0

    logger = configure_logging(runtime_root / "logs")
    try:
        with BrowserManager(
            settings.browser,
            ignore_https_errors=args.ignore_https_errors,
            stealth_enabled=(
                settings.captcha.stealth_enabled
                and is_allowed_host_url(settings.site.login_url, settings.captcha.allowed_hosts)
            ),
        ) as browser:
            page = browser.page
            try:
                print("阶段 1/2：复用现有智慧人社自动化登录")
                login = LoginService(page, settings, progress_callback=print)
                login.ensure_authenticated()
                page = login.page
                print("阶段 2/2：校验南京地区并录入人员基础信息")
                completed_rows: list[PersonInformationResult] = []

                def save_completed_row(row: PersonInformationResult) -> None:
                    completed_rows.append(row)
                    _write_results(
                        output_path,
                        PersonInformationPreparation(
                            items=tuple(items),
                            results=tuple(completed_rows),
                            city_code="320100",
                            city_name="南京",
                        ),
                    )

                result = PersonInformationService(
                    settings,
                    logger,
                    progress_callback=print,
                    result_callback=save_completed_row,
                ).prepare_with_page(page, items)
                _save_screenshot(browser.page, diagnostic_path)
                print(
                    f"采集录入完成：总计 {result.total_count} 条，"
                    f"成功 {result.prepared_count} 条，"
                    f"存在问题 {result.total_count - result.prepared_count} 条，"
                    "submitted=false"
                )
                print(f"结果 Excel：{output_path.resolve()}")
                if args.pause_after_fill:
                    input("页面停在提交前，检查完毕后按回车关闭浏览器：")
                return 0
            except (EhrmError, OSError, PlaywrightError, ValueError):
                _save_screenshot(browser.page, diagnostic_path)
                raise
    except (EhrmError, OSError, PlaywrightError, ValueError) as exc:
        _print_failure("完整人员基础信息采集 E2E 失败", exc)
        return 1


def _write_results(path: Path, result: PersonInformationPreparation) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "录入结果"
    sheet.append(
        ["Excel行号", "身份证号", "姓名", "结果", "问题编码", "问题信息", "是否提交"]
    )
    for row in result.results:
        sheet.append([
            row.item.source_index,
            row.item.identity_number,
            row.item.name,
            "已录入" if row.success else "存在问题",
            row.code,
            row.message,
            "否",
        ])
    workbook.save(path)
    workbook.close()


def _save_screenshot(page: object, path: Path) -> None:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        page.screenshot(path=str(path), full_page=True)
        print(f"页面截图：{path.resolve()}")
    except (AttributeError, OSError, PlaywrightError) as exc:
        print(f"页面截图保存失败：{exc}")


def _print_failure(label: str, exc: Exception) -> None:
    message = str(exc)
    details = getattr(exc, "details", None)
    if details and details != message:
        message += f"；{details}"
    print(f"{label}：{message}")


if __name__ == "__main__":
    raise SystemExit(main())
