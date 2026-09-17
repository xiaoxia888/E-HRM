from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries

from ehrm.core.error_catalog import ErrorCode
from ehrm.modules.employment_termination.models import (
    EmploymentTerminationItemResult,
)


_RESULT_HEADER = "处理结果"
_FAILURE_HEADER = "失败原因"
_SUCCESS_TEXT = "录入成功"
_NOT_FOUND_TEXT = "未查询到人员"
_FAILURE_TEXT = "处理失败"
_SUCCESS_FILL = PatternFill(fill_type="solid", fgColor="E8F5E9")
_FAILURE_FILL = PatternFill(fill_type="solid", fgColor="FDE9E7")


class EmploymentTerminationResultWorkbookWriter:
    """Copies the input workbook and writes one result for every source row."""

    def write(
        self,
        source: Path,
        output_dir: Path,
        results: tuple[EmploymentTerminationItemResult, ...]
        | list[EmploymentTerminationItemResult],
        *,
        destination: Path | None = None,
    ) -> Path:
        keep_vba = source.suffix.lower() == ".xlsm"
        workbook = load_workbook(source, keep_vba=keep_vba)
        try:
            sheet = workbook.active
            old_max_column = sheet.max_column
            result_column = self._column(sheet, _RESULT_HEADER)
            self._style_header(
                sheet,
                result_column,
                old_max_column,
                _RESULT_HEADER,
            )
            failure_column = self._column(sheet, _FAILURE_HEADER)
            self._style_header(
                sheet,
                failure_column,
                old_max_column,
                _FAILURE_HEADER,
            )

            by_row = {result.source_index: result for result in results}
            for row_number in range(2, sheet.max_row + 1):
                self._copy_neighbor_style(
                    sheet,
                    row_number,
                    result_column,
                    old_max_column,
                )
                self._copy_neighbor_style(
                    sheet,
                    row_number,
                    failure_column,
                    old_max_column,
                )
                status_cell = sheet.cell(row=row_number, column=result_column)
                failure_cell = sheet.cell(row=row_number, column=failure_column)
                result = by_row.get(row_number)
                if result is None:
                    status_cell.value = ""
                    failure_cell.value = ""
                    continue

                if result.success:
                    status_cell.value = _SUCCESS_TEXT
                    status_cell.fill = copy(_SUCCESS_FILL)
                    failure_cell.value = ""
                else:
                    status_cell.value = self._failure_status(result.code)
                    status_cell.fill = copy(_FAILURE_FILL)
                    failure_cell.value = result.message
                    failure_cell.fill = copy(_FAILURE_FILL)
                    failure_cell.alignment = Alignment(
                        horizontal="left",
                        vertical="center",
                        wrap_text=True,
                    )

            sheet.column_dimensions[get_column_letter(result_column)].width = 18
            sheet.column_dimensions[get_column_letter(failure_column)].width = 46
            self._extend_filters_and_tables(
                sheet,
                old_max_column,
                max(result_column, failure_column),
            )

            output_dir.mkdir(parents=True, exist_ok=True)
            resolved_destination = destination or self._destination(
                source,
                output_dir,
            )
            resolved_destination.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(resolved_destination)
            return resolved_destination
        finally:
            workbook.close()

    @staticmethod
    def _failure_status(code: str) -> str:
        if code == ErrorCode.EMPLOYEE_NOT_FOUND:
            return _NOT_FOUND_TEXT
        return _FAILURE_TEXT

    @staticmethod
    def _column(sheet, header: str) -> int:
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(row=1, column=column).value
            if str(value).strip() == header:
                return column
        return sheet.max_column + 1

    @staticmethod
    def _style_header(
        sheet,
        result_column: int,
        old_max_column: int,
        header: str,
    ) -> None:
        cell = sheet.cell(row=1, column=result_column)
        if old_max_column > 0 and result_column > old_max_column:
            source = sheet.cell(row=1, column=old_max_column)
            cell._style = copy(source._style)
            cell.font = copy(source.font)
            cell.fill = copy(source.fill)
            cell.border = copy(source.border)
            cell.alignment = copy(source.alignment)
            cell.protection = copy(source.protection)
        cell.value = header

    @staticmethod
    def _copy_neighbor_style(
        sheet,
        row_number: int,
        result_column: int,
        old_max_column: int,
    ) -> None:
        if result_column <= old_max_column or old_max_column < 1:
            return
        source = sheet.cell(row=row_number, column=old_max_column)
        target = sheet.cell(row=row_number, column=result_column)
        target._style = copy(source._style)
        target.number_format = "@"

    @staticmethod
    def _extend_filters_and_tables(
        sheet,
        old_max_column: int,
        last_result_column: int,
    ) -> None:
        new_letter = get_column_letter(last_result_column)
        if sheet.auto_filter.ref:
            min_col, min_row, max_col, max_row = range_boundaries(
                sheet.auto_filter.ref
            )
            if max_col == old_max_column:
                sheet.auto_filter.ref = (
                    f"{get_column_letter(min_col)}{min_row}:"
                    f"{new_letter}{max_row}"
                )
        for table in sheet.tables.values():
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            if max_col == old_max_column:
                table.ref = (
                    f"{get_column_letter(min_col)}{min_row}:"
                    f"{new_letter}{max_row}"
                )

    @staticmethod
    def _destination(source: Path, output_dir: Path) -> Path:
        suffix = source.suffix.lower()
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        candidate = output_dir / f"{source.stem}_退保录入结果_{stamp}{suffix}"
        if not candidate.exists():
            return candidate
        return output_dir / (
            f"{source.stem}_退保录入结果_"
            f"{datetime.now():%Y%m%d_%H%M%S_%f}{suffix}"
        )
