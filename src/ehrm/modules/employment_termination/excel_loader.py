from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ehrm.core.exceptions import ExcelValidationError, QueryValidationError
from ehrm.modules.employment_termination.models import EmploymentTerminationItem


class EmploymentTerminationExcelLoader:
    """Test adapter that converts a two-column workbook to domain objects."""

    REQUIRED_HEADERS = ("身份证号", "退工原因")

    def load(self, path: Path) -> list[EmploymentTerminationItem]:
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ExcelValidationError("退保测试文件仅支持 .xlsx 或 .xlsm")
        if not path.is_file():
            raise ExcelValidationError(f"退保测试文件不存在：{path}")
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise ExcelValidationError(
                "无法读取退保测试 Excel", details=str(exc)
            ) from exc

        try:
            rows = workbook.active.iter_rows(values_only=True)
            try:
                raw_headers = next(rows)
            except StopIteration as exc:
                raise ExcelValidationError("退保测试 Excel 为空") from exc
            headers = [
                str(value).strip() if value is not None else ""
                for value in raw_headers
            ]
            missing = [name for name in self.REQUIRED_HEADERS if name not in headers]
            if missing:
                raise ExcelValidationError(
                    "退保测试 Excel 缺少必要列：" + "、".join(missing)
                )
            identity_index = headers.index("身份证号")
            reason_index = headers.index("退工原因")
            items: list[EmploymentTerminationItem] = []
            errors: list[str] = []
            for row_number, values in enumerate(rows, start=2):
                identity_value = self._value(values, identity_index)
                reason_value = self._value(values, reason_index)
                if identity_value is None and reason_value is None:
                    continue
                try:
                    items.append(
                        EmploymentTerminationItem(
                            identity_number=self._text(identity_value),
                            termination_reason=self._text(reason_value),
                            source_index=row_number,
                        ).normalized()
                    )
                except QueryValidationError as exc:
                    errors.append(str(exc))
            if errors:
                raise ExcelValidationError(
                    "退保测试 Excel 数据校验失败",
                    details="\n".join(errors),
                )
            if not items:
                raise ExcelValidationError("退保测试 Excel 中没有可执行的数据")
            return items
        finally:
            workbook.close()

    @staticmethod
    def _value(values: tuple[object, ...], index: int) -> object | None:
        return values[index] if index < len(values) else None

    @staticmethod
    def _text(value: object | None) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()
