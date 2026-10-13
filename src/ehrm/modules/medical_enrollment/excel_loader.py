from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from ehrm.core.exceptions import ExcelValidationError, QueryValidationError
from ehrm.modules.medical_enrollment.models import MedicalPreSubmitItem, MedicalStage


class MedicalPreSubmitExcelLoader:
    """One person per run: a collection preview is not a saved archive."""

    def load(self, path: Path, stage: MedicalStage) -> MedicalPreSubmitItem:
        if path.suffix.lower() not in {".xlsx", ".xlsm"} or not path.is_file():
            raise ExcelValidationError("医保测试文件必须是存在的 .xlsx/.xlsm 文件")
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise ExcelValidationError("无法读取医保测试 Excel", details=str(exc)) from exc
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows, ())]
            required = (
                ("身份证号", "姓名") if stage is MedicalStage.COLLECTION
                else ("身份证号", "参保日期", "联系电话")
            )
            missing = [name for name in required if name not in headers]
            if missing:
                raise ExcelValidationError("医保测试 Excel 缺少列：" + "、".join(missing))
            values = [row for row in rows if any(value is not None for value in row)]
            if len(values) != 1:
                raise ExcelValidationError("医保填表测试每次仅允许一名人员")
            row = values[0]

            def value(name: str) -> object | None:
                if name not in headers or headers.index(name) >= len(row):
                    return None
                return row[headers.index(name)]

            raw_identity = value("身份证号")
            if not isinstance(raw_identity, str):
                raise ExcelValidationError(
                    "身份证号单元格必须设置为文本；Excel 数值会丢失证件号精度"
                )
            try:
                return MedicalPreSubmitItem(
                    identity_number=raw_identity,
                    name=_text(value("姓名")),
                    mobile=_text(value("联系电话")),
                    nation=_text(value("民族")),
                    document_start_date=_date(value("证件有效期起始日期"), "证件有效期起始日期"),
                    document_end_date=_date(value("证件有效期截止日期"), "证件有效期截止日期"),
                    enrollment_date=_date(value("参保日期"), "参保日期"),
                ).normalized(stage)
            except QueryValidationError as exc:
                raise ExcelValidationError(exc.message, details=exc.details) from exc
        finally:
            workbook.close()


def _text(value: object | None) -> str:
    return "" if value is None else str(value).strip()


def _date(value: object | None, label: str) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise ExcelValidationError(f"{label}必须是 YYYY-MM-DD 或 Excel 日期") from exc
