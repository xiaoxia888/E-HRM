from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from ehrm.core.exceptions import ExcelValidationError, QueryValidationError
from ehrm.modules.medical_termination.models import MedicalTerminationItem


class MedicalTerminationExcelLoader:
    REQUIRED_HEADERS = ("身份证号", "停保原因", "停保日期", "停保当月缴费")

    def load(self, path: Path) -> MedicalTerminationItem:
        if path.suffix.lower() not in {".xlsx", ".xlsm"} or not path.is_file():
            raise ExcelValidationError("医保停保测试文件必须是存在的 .xlsx/.xlsm 文件")
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise ExcelValidationError("无法读取医保停保测试 Excel", details=str(exc)) from exc
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows, ())]
            missing = [name for name in self.REQUIRED_HEADERS if name not in headers]
            if missing:
                raise ExcelValidationError("医保停保测试 Excel 缺少列：" + "、".join(missing))
            values = [row for row in rows if any(value is not None for value in row)]
            if len(values) != 1:
                raise ExcelValidationError(
                    "医保停保填表测试每次仅允许一名人员；请将 Excel 保留一条数据"
                )
            row = values[0]
            fields = {
                name: row[headers.index(name)] if headers.index(name) < len(row) else None
                for name in self.REQUIRED_HEADERS
            }
            if not isinstance(fields["身份证号"], str):
                raise ExcelValidationError(
                    "身份证号单元格必须设置为文本；Excel 数值会丢失 18 位证件号精度"
                )
            identity = self._text(fields["身份证号"])
            reason = self._text(fields["停保原因"])
            raw_date = fields["停保日期"]
            if isinstance(raw_date, datetime):
                termination_date = raw_date.date()
            elif isinstance(raw_date, date):
                termination_date = raw_date
            else:
                try:
                    termination_date = date.fromisoformat(self._text(raw_date))
                except ValueError as exc:
                    raise ExcelValidationError("停保日期必须是 YYYY-MM-DD 或 Excel 日期") from exc
            pay_text = self._text(fields["停保当月缴费"])
            if pay_text not in {"是", "否"}:
                raise ExcelValidationError("停保当月缴费只能填写“是”或“否”")
            try:
                return MedicalTerminationItem(
                    identity,
                    reason,
                    termination_date,
                    pay_text == "是",
                    source_index=2,
                ).normalized()
            except QueryValidationError as exc:
                raise ExcelValidationError(exc.message, details=exc.details) from exc
        finally:
            workbook.close()

    @staticmethod
    def _text(value: object | None) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()
