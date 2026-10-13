from datetime import date
from itertools import count
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, Mock, call, patch

from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright
import pytest

from ehrm.core.auth_repository import AuthenticationRepository, SystemType
from ehrm.core.exceptions import ExcelValidationError, QueryResultTimeoutError, WebsiteStructureChangedError
from ehrm.core.preferences import UserPreferences
from ehrm.core.settings import BrowserInteractionPacingSettings, load_settings
from ehrm.entrypoints.medical_enrollment_e2e_cli import _write_result, main
from ehrm.modules.medical_enrollment.excel_loader import MedicalPreSubmitExcelLoader
from ehrm.modules.medical_enrollment.models import MedicalPreSubmitItem, MedicalStage
from ehrm.modules.medical_enrollment.page import (
    MedicalPreSubmitBlockedError,
    MedicalPreSubmitPage,
    identity_demographics,
)
from ehrm.modules.medical_enrollment.submission import (
    MedicalEnrollmentSubmissionPage,
    MedicalEnrollmentSubmissionUncertainError,
    classify_collection_response,
    confirm_enrollment_response,
)
from ehrm.web.settings_service import WebSettingsService


_PORTAL = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"


@pytest.fixture(scope="module")
def enrollment_local_page():
    """Only a local replica of the supplied DOM; never contacts the portal."""
    with sync_playwright() as playwright:
        if not Path(playwright.chromium.executable_path).exists():
            pytest.skip("本地 DOM 回归测试需要已安装的 Playwright Chromium")
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page()
        yield page
        browser.close()


def _set_enrollment_dom(page) -> None:
    page.set_content('''
        <form name="f_form">
          <div class="el-card">
            <div class="v_header_title_content">员工个人信息查询</div>
            <div class="el-form-item"><label for="certno">身份证号码(输入后点击查询)：</label>
              <div class="el-form-item__content"><input class="el-input__inner"></div></div>
            <button type="button" data-test="person-query"
              onclick="window.personQueries=(window.personQueries||0)+1">查询</button>
            <div class="el-form-item"><label for="psnName">姓名：</label>
              <div class="el-form-item__content"><input class="el-input__inner" disabled value="测试人员"></div></div>
          </div>
          <div class="el-card">
            <div class="v_header_title_content">单位员工参保信息</div>
            <div class="el-form-item"><label for="zjhm">证件号码</label>
              <div class="el-form-item__content"><input class="el-input__inner" disabled value="320101199001011234"></div></div>
            <div class="el-form-item"><label for="ddwsj">到单位时间</label>
              <div class="el-form-item__content"><input class="el-input__inner" placeholder="选择日期"></div></div>
            <div class="el-form-item"><label for="mob">手机号</label>
              <div class="el-form-item__content"><input class="el-input__inner"></div></div>
            <button type="button" data-test="enrollment-submit">提交</button>
          </div>
          <div class="el-card">
            <div class="v_header_title_content">查询办理结果</div>
            <div class="el-form-item"><label for="starTime">办理时间</label>
              <div class="el-form-item__content"><input class="el-input__inner" placeholder="选择日期"></div></div>
            <button type="button" data-test="result-query"
              onclick="window.resultQueries=(window.resultQueries||0)+1">查询</button>
          </div>
        </form>
    ''')
    page.evaluate("window.personQueries = 0; window.resultQueries = 0")


def test_enrollment_dom_uses_field_labels_and_correct_query_without_submitting(
    enrollment_local_page,
) -> None:
    _set_enrollment_dom(enrollment_local_page)
    with patch("ehrm.modules.medical_enrollment.page.is_medical_portal_url", return_value=True):
        automation = MedicalPreSubmitPage(enrollment_local_page, MedicalStage.ENROLLMENT)
    automation._wait_for_query = Mock()  # type: ignore[method-assign]
    item = MedicalPreSubmitItem(
        "320101199001011234", name="测试人员", mobile="13800138000",
        enrollment_date=date(2026, 10, 12),
    )

    automation._fill_enrollment(item)

    assert enrollment_local_page.evaluate("window.personQueries") == 1
    assert enrollment_local_page.evaluate("window.resultQueries || 0") == 0
    assert automation._enrollment_identity().input_value() == item.identity_number
    assert automation._enrollment_input("ddwsj", "到单位时间").input_value() == "2026-10-12"
    assert automation._enrollment_input("mob", "手机号").input_value() == item.mobile
    assert automation._enrollment_input("psnName", "姓名", require_enabled=False).input_value() == item.name
    assert enrollment_local_page.locator('[data-test="enrollment-submit"]').is_enabled()


def test_enrollment_open_reuses_already_visible_form(enrollment_local_page) -> None:
    _set_enrollment_dom(enrollment_local_page)
    progress = Mock()
    with patch("ehrm.modules.medical_enrollment.page.is_medical_portal_url", return_value=True):
        automation = MedicalPreSubmitPage(
            enrollment_local_page, MedicalStage.ENROLLMENT, progress_callback=progress
        )

    automation.open()
    automation.open()

    assert progress.call_args_list == [
        call("医保：人员参保登记已打开，复用当前页面"),
        call("医保：人员参保登记已打开，复用当前页面"),
    ]
    assert automation._enrollment_identity().input_value() == ""


def test_enrollment_can_fill_second_person_on_same_open_page(enrollment_local_page) -> None:
    _set_enrollment_dom(enrollment_local_page)
    with patch("ehrm.modules.medical_enrollment.page.is_medical_portal_url", return_value=True):
        automation = MedicalPreSubmitPage(enrollment_local_page, MedicalStage.ENROLLMENT)
    automation._wait_for_query = Mock()  # type: ignore[method-assign]
    first = MedicalPreSubmitItem(
        "320101199001011234", name="测试人员", mobile="13800138000",
        enrollment_date=date(2026, 10, 12),
    )
    second = MedicalPreSubmitItem(
        "320101199001011235", name="另一人员", mobile="13800138001",
        enrollment_date=date(2026, 10, 13),
    )

    automation.fill_item(first)
    automation._enrollment_identity().fill("")
    enrollment_local_page.locator('label[for="zjhm"] + .el-form-item__content input').evaluate(
        "element => element.value = '320101199001011235'"
    )
    enrollment_local_page.locator('label[for="psnName"] + .el-form-item__content input').evaluate(
        "element => element.value = '另一人员'"
    )
    automation.fill_item(second)

    assert enrollment_local_page.evaluate("window.personQueries") == 2
    assert enrollment_local_page.evaluate("window.resultQueries || 0") == 0
    assert automation._enrollment_identity().input_value() == second.identity_number
    assert automation._enrollment_input("ddwsj", "到单位时间").input_value() == "2026-10-13"
    assert automation._enrollment_input("mob", "手机号").input_value() == second.mobile


def test_enrollment_submit_checks_same_fields_without_clicking_in_local_dom(
    enrollment_local_page,
) -> None:
    _set_enrollment_dom(enrollment_local_page)
    with patch("ehrm.modules.medical_enrollment.page.is_medical_portal_url", return_value=True):
        automation = MedicalEnrollmentSubmissionPage(enrollment_local_page, MedicalStage.ENROLLMENT)
    item = MedicalPreSubmitItem(
        "320101199001011234", name="测试人员", mobile="13800138000",
        enrollment_date=date(2026, 10, 12),
    )
    automation._wait_for_query = Mock()  # type: ignore[method-assign]
    automation._fill_enrollment(item)
    automation._click_and_wait_response = Mock(return_value={  # type: ignore[method-assign]
        "code": 0, "type": "success", "message": "办理成功!",
    })
    automation._has_result_row = Mock(return_value=True)  # type: ignore[method-assign]
    automation._wait_for_result_row = Mock()  # type: ignore[method-assign]

    assert "办理成功" in automation.submit_enrollment(item, before_submit=Mock())
    button = automation._click_and_wait_response.call_args.args[0]
    assert button.get_attribute("data-test") == "enrollment-submit"
    assert automation._click_and_wait_response.call_count == 1


def test_enrollment_submit_rejects_other_person_before_submit(enrollment_local_page) -> None:
    _set_enrollment_dom(enrollment_local_page)
    with patch("ehrm.modules.medical_enrollment.page.is_medical_portal_url", return_value=True):
        automation = MedicalEnrollmentSubmissionPage(enrollment_local_page, MedicalStage.ENROLLMENT)
    item = MedicalPreSubmitItem(
        "320101199001011234", name="测试人员", mobile="13800138000",
        enrollment_date=date(2026, 10, 12),
    )
    automation._wait_for_query = Mock()  # type: ignore[method-assign]
    automation._fill_enrollment(item)
    enrollment_local_page.locator('label[for="zjhm"] + .el-form-item__content input').evaluate(
        "element => element.value = '320101199001011235'"
    )
    automation._click_and_wait_response = Mock()  # type: ignore[method-assign]

    with pytest.raises(WebsiteStructureChangedError, match="证件号码与申请不一致"):
        automation.submit_enrollment(item, before_submit=Mock())
    automation._click_and_wait_response.assert_not_called()


def test_collection_demographics_are_derived_from_identity() -> None:
    assert identity_demographics('320101199001011234') == (date(1990, 1, 1), '男')
    assert identity_demographics('320101199001011224') == (date(1990, 1, 1), '女')
    with pytest.raises(WebsiteStructureChangedError, match='出生日期'):
        identity_demographics('320101199013011234')


def test_recorded_collection_duplicate_is_existing_not_failed() -> None:
    assert classify_collection_response({
        'code': -1, 'type': 'error', 'message': '该人员已经建档,请勿重复建档!', 'data': None,
    }) == 'existing'


def test_recorded_enrollment_success_requires_explicit_business_confirmation() -> None:
    assert '办理成功' in confirm_enrollment_response({
        'code': 0, 'type': 'success', 'message': '办理成功!,请于第二天再登录税务系统查询（申报期内）!',
    })
    with pytest.raises(MedicalEnrollmentSubmissionUncertainError):
        confirm_enrollment_response({'code': 0, 'type': 'success', 'message': '已收到申请'})
    with pytest.raises(MedicalEnrollmentSubmissionUncertainError):
        classify_collection_response({'code': 0, 'type': 'success', 'message': '已收到申请'})


def test_submit_listener_uses_exact_medical_endpoint_and_clicks_once() -> None:
    listeners = {}
    response = SimpleNamespace(
        url='https://wb.njybjyybz.org.cn:9091/hsa-pss-cw-local-nj/api/dwctbbl/rycbdj/saveInfo',
        request=SimpleNamespace(method='POST'), status=200,
        json=lambda: {'code': 0, 'type': 'success', 'message': '办理成功!'},
    )
    page = SimpleNamespace(
        url=_PORTAL,
        on=lambda event, callback: listeners.setdefault(event, callback),
        remove_listener=lambda event, callback: listeners.pop(event),
        wait_for_timeout=Mock(),
    )
    button = SimpleNamespace(click=Mock(side_effect=lambda: listeners['response'](response)))
    before = Mock()
    automation = MedicalEnrollmentSubmissionPage(page, MedicalStage.ENROLLMENT)  # type: ignore[arg-type]
    assert automation._click_and_wait_response(button, '/api/dwctbbl/rycbdj/saveInfo', before) == response.json()
    before.assert_called_once()
    button.click.assert_called_once()
    assert listeners == {}


def test_submit_listener_treats_missing_response_as_uncertain_without_second_click() -> None:
    page = SimpleNamespace(
        url=_PORTAL, on=Mock(), remove_listener=Mock(), wait_for_timeout=Mock(),
    )
    button = SimpleNamespace(click=Mock())
    before = Mock()
    automation = MedicalEnrollmentSubmissionPage(page, MedicalStage.ENROLLMENT, timeout_ms=1)  # type: ignore[arg-type]
    with pytest.raises(MedicalEnrollmentSubmissionUncertainError, match='请先到医保平台核对'):
        automation._click_and_wait_response(button, '/api/dwctbbl/rycbdj/saveInfo', before)
    before.assert_called_once()
    button.click.assert_called_once()


def test_enrollment_submits_once_then_queries_result_without_resubmitting() -> None:
    identity = SimpleNamespace(input_value=lambda: '320101199001011234')
    date_field = SimpleNamespace(input_value=lambda: '2026-10-12')
    mobile = SimpleNamespace(input_value=lambda: '13800138000')
    name = SimpleNamespace(input_value=lambda: '测试人员')
    displayed_identity = SimpleNamespace(input_value=lambda: '320101199001011234')
    submit_button = object()
    result_button = object()
    result_card = SimpleNamespace(get_by_role=Mock(return_value=result_button))
    cards = SimpleNamespace(filter=Mock(return_value=object()))
    page = SimpleNamespace(
        url=_PORTAL, locator=Mock(return_value=cards),
    )
    automation = MedicalEnrollmentSubmissionPage(page, MedicalStage.ENROLLMENT)  # type: ignore[arg-type]
    automation._enrollment_identity = Mock(return_value=identity)  # type: ignore[method-assign]
    fields = {'zjhm': displayed_identity, 'psnName': name, 'ddwsj': date_field, 'mob': mobile}
    automation._enrollment_input = Mock(side_effect=lambda field_id, _label, **_: fields[field_id])  # type: ignore[method-assign]
    automation._enrollment_button = Mock(return_value=submit_button)  # type: ignore[method-assign]
    automation._visible_nodes = Mock(return_value=[result_card])  # type: ignore[method-assign]
    automation._unique_button = Mock(side_effect=lambda _nodes, label: submit_button if label == '医保参保提交' else result_button)  # type: ignore[method-assign]
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    automation._has_result_row = Mock(return_value=False)  # type: ignore[method-assign]
    automation._wait_for_result_row = Mock()  # type: ignore[method-assign]
    automation._click_and_wait_response = Mock(side_effect=[  # type: ignore[method-assign]
        {'code': 0, 'type': 'success', 'message': '办理成功!'},
        {'code': 0, 'type': 'success', 'message': '成功'},
    ])
    before = Mock()
    item = MedicalPreSubmitItem('320101199001011234', name='测试人员', mobile='13800138000',
                                enrollment_date=date(2026, 10, 12))

    assert '办理成功' in automation.submit_enrollment(item, before_submit=before)
    assert automation._click_and_wait_response.call_args_list == [
        call(submit_button, '/api/dwctbbl/rycbdj/saveInfo', before),
        call(result_button, '/api/dwctbbl/rycbdj/queryBl'),
    ]
    automation._enrollment_button.assert_called_once_with('单位员工参保信息', '提交')
    automation._wait_for_result_row.assert_called_once_with(item.identity_number)


def _workbook(path: Path, rows: list[list[object]]) -> Path:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    workbook.close()
    return path


def _fake_page(message: str) -> SimpleNamespace:
    node = SimpleNamespace(is_visible=lambda: True, inner_text=lambda: message)
    nodes = SimpleNamespace(count=lambda: 1, nth=lambda _index: node)
    empty = SimpleNamespace(count=lambda: 0)
    return SimpleNamespace(
        url=_PORTAL,
        locator=lambda selector: (
            nodes if selector == ".el-message-box__message" else empty
        ),
    )


def test_collection_excel_loads_recorded_fields(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "采集.xlsx", [
        ["身份证号", "姓名", "联系电话", "民族", "证件有效期起始日期", "证件有效期截止日期"],
        ["320101199001011234", "测试人员", "13800138000", "汉族", "2020-01-01", "2040-01-01"],
    ])
    item = MedicalPreSubmitExcelLoader().load(path, MedicalStage.COLLECTION)
    assert item.name == "测试人员"
    assert item.document_start_date == date(2020, 1, 1)
    assert item.document_end_date == date(2040, 1, 1)


def test_enrollment_requires_date_and_valid_mobile(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "参保.xlsx", [
        ["身份证号", "参保日期", "联系电话"],
        ["320101199001011234", "2026-10-12", "12345"],
    ])
    with pytest.raises(ExcelValidationError, match="11 位"):
        MedicalPreSubmitExcelLoader().load(path, MedicalStage.ENROLLMENT)


def test_excel_requires_text_identity_and_one_person(tmp_path: Path) -> None:
    numeric = _workbook(tmp_path / "证件号数字.xlsx", [
        ["身份证号", "姓名"],
        [320101199001011234, "测试人员"],
    ])
    with pytest.raises(ExcelValidationError, match="必须设置为文本"):
        MedicalPreSubmitExcelLoader().load(numeric, MedicalStage.COLLECTION)
    two = _workbook(tmp_path / "两人.xlsx", [
        ["身份证号", "姓名"],
        ["320101199001011234", "甲"],
        ["320101199001011235", "乙"],
    ])
    with pytest.raises(ExcelValidationError, match="仅允许一名人员"):
        MedicalPreSubmitExcelLoader().load(two, MedicalStage.COLLECTION)


@pytest.mark.parametrize("message,outcome", [
    ("请先给该人员建档!", "需先建档"),
    ("该人员已经建档,请勿重复建档!", "已建档"),
    ("该人员已经参保!请勿重复参保!!", "已参保"),
    ("请输入正确的手机号", "业务阻断"),
])
def test_portal_feedback_is_blocking_not_success(message: str, outcome: str) -> None:
    automation = MedicalPreSubmitPage(_fake_page(message), MedicalStage.ENROLLMENT)  # type: ignore[arg-type]
    with pytest.raises(MedicalPreSubmitBlockedError) as raised:
        automation._raise_feedback()
    assert raised.value.outcome == outcome


def test_query_waits_for_loading_to_end() -> None:
    page = SimpleNamespace(url=_PORTAL, wait_for_timeout=Mock())
    automation = MedicalPreSubmitPage(page, MedicalStage.ENROLLMENT)  # type: ignore[arg-type]
    states = iter([True, True, False, False, False])
    automation._any_visible = Mock(  # type: ignore[method-assign]
        side_effect=lambda _selector: next(states)
    )
    automation._enrollment_input = Mock(side_effect=lambda field_id, _label, **_:  # type: ignore[method-assign]
        SimpleNamespace(input_value=lambda: '320101199001011234' if field_id == 'zjhm' else '测试人员'))
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    clock = count(step=0.2)
    response = SimpleNamespace(status=200, json=lambda: {'code': 0, 'type': 'success'})
    with patch("ehrm.modules.medical_enrollment.page.time.monotonic", side_effect=lambda: next(clock)):
        automation._wait_for_query([response], '320101199001011234')  # type: ignore[list-item]
    assert page.wait_for_timeout.call_count >= 3


def test_query_api_business_error_is_not_a_filled_form() -> None:
    page = SimpleNamespace(url=_PORTAL, wait_for_timeout=Mock())
    automation = MedicalPreSubmitPage(page, MedicalStage.ENROLLMENT)  # type: ignore[arg-type]
    automation._any_visible = Mock(return_value=False)  # type: ignore[method-assign]
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    response = SimpleNamespace(
        status=200,
        json=lambda: {"code": 1, "message": "请先给该人员建档!"},
    )
    clock = count(step=0.2)
    with patch("ehrm.modules.medical_enrollment.page.time.monotonic", side_effect=lambda: next(clock)):
        with pytest.raises(MedicalPreSubmitBlockedError) as raised:
            automation._wait_for_query([response], '320101199001011234')  # type: ignore[list-item]
    assert raised.value.outcome == "需先建档"


def test_query_explicit_ok_is_not_overridden_by_nonzero_code() -> None:
    page = SimpleNamespace(url=_PORTAL, wait_for_timeout=Mock())
    automation = MedicalPreSubmitPage(page, MedicalStage.ENROLLMENT)  # type: ignore[arg-type]
    automation._any_visible = Mock(return_value=False)  # type: ignore[method-assign]
    automation._enrollment_input = Mock(side_effect=lambda field_id, _label, **_:  # type: ignore[method-assign]
        SimpleNamespace(input_value=lambda: '320101199001011234' if field_id == 'zjhm' else '测试人员'))
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    response = SimpleNamespace(
        status=200,
        json=lambda: {"code": 1, "ok": True, "errmsg": "未提供错误信息"},
    )
    clock = count(step=0.2)
    with patch("ehrm.modules.medical_enrollment.page.time.monotonic", side_effect=lambda: next(clock)):
        automation._wait_for_query([response], '320101199001011234')  # type: ignore[list-item]


def test_query_does_not_treat_result_search_date_as_person_ready() -> None:
    page = SimpleNamespace(url=_PORTAL, wait_for_timeout=Mock())
    automation = MedicalPreSubmitPage(page, MedicalStage.ENROLLMENT)  # type: ignore[arg-type]
    automation._any_visible = Mock(return_value=False)  # type: ignore[method-assign]
    automation._enrollment_input = Mock(side_effect=lambda field_id, _label, **_:  # type: ignore[method-assign]
        SimpleNamespace(input_value=lambda: '320101199001011235' if field_id == 'zjhm' else '其他人员'))
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    response = SimpleNamespace(status=200, json=lambda: {'code': 0, 'type': 'success'})
    clock = count(step=5)

    with patch("ehrm.modules.medical_enrollment.page.time.monotonic", side_effect=lambda: next(clock)):
        with pytest.raises(WebsiteStructureChangedError, match='未显示当前身份证号及姓名'):
            automation._wait_for_query([response], '320101199001011234')  # type: ignore[list-item]


def test_query_without_observed_api_response_stops_before_filling() -> None:
    page = SimpleNamespace(url=_PORTAL, wait_for_timeout=Mock())
    automation = MedicalPreSubmitPage(page, MedicalStage.ENROLLMENT)  # type: ignore[arg-type]
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    automation._enrollment_input = Mock()  # type: ignore[method-assign]
    clock = count(step=5)

    with patch("ehrm.modules.medical_enrollment.page.time.monotonic", side_effect=lambda: next(clock)):
        with pytest.raises(QueryResultTimeoutError, match='查询加载超时'):
            automation._wait_for_query([], '320101199001011234')
    automation._enrollment_input.assert_not_called()


def test_recorded_identity_position_must_match_label() -> None:
    label = SimpleNamespace(inner_text=lambda: "姓名：")
    form_item = SimpleNamespace(locator=lambda _selector: SimpleNamespace(first=label))
    field = SimpleNamespace(locator=lambda _selector: form_item)
    automation = MedicalPreSubmitPage(SimpleNamespace(url=_PORTAL), MedicalStage.COLLECTION)  # type: ignore[arg-type]
    with pytest.raises(WebsiteStructureChangedError, match="顺序已变化"):
        automation._assert_context(field, "证件")  # type: ignore[arg-type]


def test_collection_identity_uses_certno_field_not_textbox_position() -> None:
    identity = object()
    automation = MedicalPreSubmitPage(SimpleNamespace(url=_PORTAL), MedicalStage.COLLECTION)  # type: ignore[arg-type]
    automation._collection_input = Mock(return_value=identity)  # type: ignore[method-assign]

    assert automation._collection_identity() is identity
    automation._collection_input.assert_called_once_with("certno", "证件号码")


@pytest.mark.parametrize(("field_id", "label_text"), [
    ("certno", "证件号码"),
    ("psnName", "姓名"),
    ("mob", "联系电话"),
    ("naty", "民族"),
    ("zjyxqqsrq", "证件有效期(起始日期)"),
    ("zjyxqjzrq", "证件有效期(截止日期)"),
])
def test_collection_input_uses_form_label_for_and_adjacent_content(
    field_id: str, label_text: str,
) -> None:
    label = SimpleNamespace(
        first=SimpleNamespace(wait_for=Mock()),
        count=lambda: 1,
        inner_text=lambda: f"{label_text}：",
    )
    field = SimpleNamespace(
        first=SimpleNamespace(wait_for=Mock()),
        count=lambda: 1,
        is_enabled=lambda: True,
    )
    form = SimpleNamespace(locator=Mock(side_effect=[label, field]))
    page = SimpleNamespace(url=_PORTAL, locator=Mock(return_value=form))
    automation = MedicalPreSubmitPage(page, MedicalStage.COLLECTION)  # type: ignore[arg-type]

    assert automation._collection_input(field_id, label_text) is field
    page.locator.assert_called_once_with('form[name="f_form"]')
    assert [call.args[0] for call in form.locator.call_args_list] == [
        f'label[for="{field_id}"]',
        f'label[for="{field_id}"] + .el-form-item__content input.el-input__inner',
    ]


def test_collection_input_rejects_mismatched_label() -> None:
    label = SimpleNamespace(
        first=SimpleNamespace(wait_for=Mock()),
        count=lambda: 1,
        inner_text=lambda: "其他字段：",
    )
    form = SimpleNamespace(locator=Mock(return_value=label))
    page = SimpleNamespace(url=_PORTAL, locator=Mock(return_value=form))
    automation = MedicalPreSubmitPage(page, MedicalStage.COLLECTION)  # type: ignore[arg-type]

    with pytest.raises(WebsiteStructureChangedError, match="标签已变化"):
        automation._collection_input("certno", "证件号码")
    assert form.locator.call_count == 1


def test_collection_fills_recorded_fields_by_form_ids_without_submitting() -> None:
    page = SimpleNamespace(url=_PORTAL, wait_for_timeout=Mock())
    automation = MedicalPreSubmitPage(page, MedicalStage.COLLECTION)  # type: ignore[arg-type]
    fields = {
        field_id: SimpleNamespace(press=Mock(), input_value=lambda: "")
        for field_id in ("certno", "psnName", "brdy", "gend", "mob", "naty", "zjyxqqsrq", "zjyxqjzrq")
    }
    automation._collection_input = Mock(side_effect=lambda field_id, _label: fields[field_id])  # type: ignore[method-assign]
    automation._fill_verified = Mock()  # type: ignore[method-assign]
    automation._wait_until_not_loading = Mock()  # type: ignore[method-assign]
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    automation._select_option = Mock()  # type: ignore[method-assign]
    automation._set_date = Mock()  # type: ignore[method-assign]
    item = MedicalPreSubmitItem(
        "320101199001011234", name="测试人员", mobile="13800138000", nation="汉族",
        document_start_date=date(2020, 1, 1), document_end_date=date(2040, 1, 1),
    )

    automation._fill_collection(item)

    assert automation._collection_input.call_args_list == [
        call("certno", "证件号码"),
        call("brdy", "出生日期"),
        call("gend", "性别"),
        call("psnName", "姓名"),
        call("mob", "联系电话"),
        call("naty", "民族"),
        call("zjyxqqsrq", "证件有效期(起始日期)"),
        call("zjyxqjzrq", "证件有效期(截止日期)"),
    ]
    fields["certno"].press.assert_called_once_with("Enter")
    assert automation._select_option.call_args_list == [
        call("性别", "男", field=fields["gend"]),
        call("民族", "汉族", field=fields["naty"]),
    ]
    assert automation._set_date.call_args_list[0] == call(fields["brdy"], date(1990, 1, 1), "出生日期")
    assert not hasattr(automation, "submit")


def test_collection_menu_and_field_actions_use_configured_pacing() -> None:
    menu = SimpleNamespace(first=SimpleNamespace(click=Mock()))
    link = SimpleNamespace(click=Mock())
    identity = SimpleNamespace(wait_for=Mock(), fill=Mock(), input_value=lambda: "测试人员")
    page = SimpleNamespace(
        url=_PORTAL,
        wait_for_timeout=Mock(),
        get_by_text=Mock(return_value=menu),
        get_by_role=Mock(return_value=link),
    )
    pacing = BrowserInteractionPacingSettings(True, 200, 200)
    automation = MedicalPreSubmitPage(  # type: ignore[arg-type]
        page, MedicalStage.COLLECTION, pacing=pacing,
    )
    automation._collection_identity = Mock(return_value=identity)  # type: ignore[method-assign]

    automation.open()
    automation._fill_verified(identity, "测试人员", "姓名")  # type: ignore[arg-type]

    menu.first.click.assert_called_once()
    link.click.assert_called_once()
    identity.fill.assert_called_once_with("测试人员")
    assert page.wait_for_timeout.call_args_list == [call(100)] * 6
    identity.wait_for.assert_called_with(state="visible", timeout=30_000)


def test_enrollment_menu_actions_use_configured_pacing() -> None:
    menu = SimpleNamespace(first=SimpleNamespace(click=Mock()))
    link = SimpleNamespace(first=SimpleNamespace(wait_for=Mock()), click=Mock())
    page = SimpleNamespace(
        url=_PORTAL,
        wait_for_timeout=Mock(),
        get_by_text=Mock(return_value=menu),
        get_by_role=Mock(return_value=link),
    )
    automation = MedicalPreSubmitPage(  # type: ignore[arg-type]
        page, MedicalStage.ENROLLMENT,
        pacing=BrowserInteractionPacingSettings(True, 100, 100),
    )
    automation._enrollment_form_is_active = Mock(return_value=False)  # type: ignore[method-assign]
    automation._visible_nodes = Mock(side_effect=[[], [link]])  # type: ignore[method-assign]
    automation._enrollment_identity = Mock()  # type: ignore[method-assign]
    automation._enrollment_button = Mock()  # type: ignore[method-assign]

    automation.open()

    menu.first.click.assert_called_once()
    link.click.assert_called_once()
    link.first.wait_for.assert_called_once_with(state="visible", timeout=30_000)
    automation._enrollment_identity.assert_called_once()
    automation._enrollment_button.assert_called_once_with('员工个人信息查询', '查询')
    assert page.wait_for_timeout.call_args_list == [call(100), call(100)]


def test_enrollment_open_does_not_collapse_already_expanded_menu() -> None:
    menu = SimpleNamespace(first=SimpleNamespace(click=Mock()))
    link = SimpleNamespace(first=SimpleNamespace(wait_for=Mock()), click=Mock())
    page = SimpleNamespace(
        url=_PORTAL, get_by_text=Mock(return_value=menu),
        get_by_role=Mock(return_value=link),
    )
    automation = MedicalPreSubmitPage(page, MedicalStage.ENROLLMENT)  # type: ignore[arg-type]
    automation._enrollment_form_is_active = Mock(return_value=False)  # type: ignore[method-assign]
    automation._visible_nodes = Mock(return_value=[link])  # type: ignore[method-assign]
    automation._enrollment_identity = Mock()  # type: ignore[method-assign]
    automation._enrollment_button = Mock()  # type: ignore[method-assign]

    automation.open()

    menu.first.click.assert_not_called()
    link.click.assert_called_once()


def test_fill_item_never_calls_save_or_submit() -> None:
    automation = MedicalPreSubmitPage(SimpleNamespace(url=_PORTAL), MedicalStage.COLLECTION)  # type: ignore[arg-type]
    automation.open = Mock()  # type: ignore[method-assign]
    automation._fill_collection = Mock()  # type: ignore[method-assign]
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    item = MedicalPreSubmitItem("320101199001011234", name="测试人员")
    automation.fill_item(item)
    automation._fill_collection.assert_called_once()
    assert not hasattr(automation, "submit")


def test_result_explicitly_records_not_submitted(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "参保.xlsx", [
        ["身份证号", "参保日期", "联系电话"],
        ["320101199001011234", "2026-10-12", "13800138000"],
    ])
    result = _write_result(source, tmp_path / "output", MedicalStage.ENROLLMENT,
                           "已填写待核对（未提交）", "还需核对")
    workbook = load_workbook(result, read_only=True)
    try:
        headers = [cell.value for cell in workbook.active[1]]
        assert workbook.active.cell(2, headers.index("是否提交") + 1).value == "否"
    finally:
        workbook.close()


def test_check_input_uses_medical_account_without_browser(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    AuthenticationRepository(settings.auth_database_path).save_account(
        SystemType.MEDICAL, "government-user", "government-password",
        profile={"login_method": "government_sso"},
    )
    source = _workbook(tmp_path / "参保.xlsx", [
        ["身份证号", "参保日期", "联系电话"],
        ["320101199001011234", "2026-10-12", "13800138000"],
    ])
    with (
        patch("ehrm.entrypoints.medical_enrollment_e2e_cli.application_runtime_root", return_value=tmp_path),
        patch("ehrm.entrypoints.medical_enrollment_e2e_cli.load_settings", return_value=settings),
        patch("ehrm.entrypoints.medical_enrollment_e2e_cli.BrowserManager") as browser,
    ):
        status = main(["--stage", "enrollment", "--input", str(source), "--check-input"])
    assert status == 0
    browser.assert_not_called()


def test_collection_cli_passes_saved_execution_speed_to_business_page(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    AuthenticationRepository(settings.auth_database_path).save_account(
        SystemType.MEDICAL, "government-user", "government-password",
        profile={"login_method": "government_sso"},
    )
    WebSettingsService(settings, tmp_path).preferences_store.save(
        UserPreferences(execution_speed="stable")
    )
    source = _workbook(tmp_path / "采集.xlsx", [
        ["身份证号", "姓名"], ["320101199001011234", "测试人员"],
    ])
    portal = SimpleNamespace(url=_PORTAL)
    browser = MagicMock()
    browser.return_value.__enter__.return_value.page = portal
    with (
        patch("ehrm.entrypoints.medical_enrollment_e2e_cli.application_runtime_root", return_value=tmp_path),
        patch("ehrm.entrypoints.medical_enrollment_e2e_cli.load_settings", return_value=settings),
        patch("ehrm.entrypoints.medical_enrollment_e2e_cli.BrowserManager", browser),
        patch("ehrm.entrypoints.medical_enrollment_e2e_cli.LoginService") as login,
        patch("ehrm.entrypoints.medical_enrollment_e2e_cli.MedicalPreSubmitPage") as business_page,
        patch("ehrm.entrypoints.medical_enrollment_e2e_cli._save_screenshot"),
        patch("ehrm.entrypoints.medical_enrollment_e2e_cli._write_result", return_value=tmp_path / "result.xlsx"),
    ):
        login.return_value.ensure_external_government_authenticated.return_value = portal
        assert main(["--stage", "collection", "--input", str(source)]) == 0

    business_page.return_value.fill_item.assert_called_once()
    configured = business_page.call_args.kwargs["pacing"]
    assert configured.enabled is True
    assert (configured.min_delay_ms, configured.max_delay_ms) == (1500, 2500)
