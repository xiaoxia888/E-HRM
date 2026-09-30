from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import MagicMock, Mock, patch

import pytest

from ehrm.core.settings import load_settings
from ehrm.modules.employment_termination.models import EmploymentTerminationItem
from ehrm.modules.employment_termination.submission import (
    TerminationSubmissionError,
    TerminationSubmissionMonitor,
)
from ehrm.modules.employment_termination.service import EmploymentTerminationService


IDENTITY = "320101199001011234"
HOST = "rs.jshrss.jiangsu.gov.cn"
APPLY_PATH = "/api/ehrss-si-enterprise-app/api/applys/employee/dimissionApply"
COMPLETE_PATH = "/api/ggfw-affair-base-center/affair/case/commbusiness/complete"
SAVE_PATH = "/api/ehrss-si-enterprise-app/api/changeAndLookController/submitSave"


def response(path: str, body: object, *, status: int = 200, payload=None):
    result = Mock(url=f"https://{HOST}{path}", status=status)
    result.request.resource_type = "xhr"
    result.request.method = "POST"
    result.request.post_data = json.dumps(payload or {})
    result.json.return_value = body
    result.text.return_value = "" if body is None else json.dumps(body)
    return result


def monitor(tmp_path: Path):
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    return TerminationSubmissionMonitor(Mock(), settings, EmploymentTerminationItem(IDENTITY, "6304"))


def accepted(m):
    payload = {"idNumber": IDENTITY, "changeReason": "6304", "applyId": "application-1"}
    m.capture(response(SAVE_PATH, None, payload=payload))
    m.capture(response(APPLY_PATH, {"ok": True, "appcode": 1, "errmsg": "未提供错误信息", "certificateNumber": IDENTITY, "partyId": "person-1"}, status=201, payload=payload))
    m.capture(response(COMPLETE_PATH, {"appcode": "0", "msg": "", "map": {"data": {"complete": False, "busiId": "application-1"}}}))


def test_real_success_shape_needs_ui_and_accepts_complete_false(tmp_path):
    m = monitor(tmp_path)
    accepted(m)
    assert m.receipt(success_visible=False) is None
    assert m.receipt(success_visible=True).application_id == "application-1"
    m.close()
    m.page.remove_listener.assert_called_once()


@pytest.mark.parametrize("body", [{"ok": False, "errmsg": "无法办理"}, {"ok": "true"}, {"appcode": 1}])
def test_refusal_or_unknown_response_is_not_success(tmp_path, body):
    m = monitor(tmp_path)
    m.capture(response(APPLY_PATH, body, payload={"idNumber": IDENTITY, "changeReason": "6304", "applyId": "application-1"}))
    with pytest.raises(TerminationSubmissionError):
        m.receipt(success_visible=True)


@pytest.mark.parametrize("field,value", [("idNumber", "320101199002021235"), ("changeReason", "6305")])
def test_current_person_and_reason_must_match(tmp_path, field, value):
    m = monitor(tmp_path)
    payload = {"idNumber": IDENTITY, "changeReason": "6304", "applyId": "application-1", field: value}
    m.capture(response(APPLY_PATH, {"ok": True}, payload=payload))
    with pytest.raises(TerminationSubmissionError) as exc:
        m.receipt(success_visible=True)
    assert exc.value.outcome == "unknown"


def test_completion_must_match_current_application(tmp_path):
    m = monitor(tmp_path)
    accepted(m)
    m.responses[COMPLETE_PATH] = [response(COMPLETE_PATH, {"appcode": "0", "map": {"data": {"busiId": "another-application"}}})]
    with pytest.raises(TerminationSubmissionError):
        m.receipt(success_visible=True)


def test_save_and_apply_must_belong_to_same_application(tmp_path):
    m = monitor(tmp_path)
    accepted(m)
    m.responses[SAVE_PATH][0].request.post_data = json.dumps({"idNumber": IDENTITY, "changeReason": "6304", "applyId": "another-application"})
    with pytest.raises(TerminationSubmissionError):
        m.receipt(success_visible=True)


def test_response_certificate_number_must_match(tmp_path):
    m = monitor(tmp_path)
    accepted(m)
    m.responses[APPLY_PATH][0].json.return_value["certificateNumber"] = "320101199002021235"
    with pytest.raises(TerminationSubmissionError):
        m.receipt(success_visible=True)


def test_save_failure_overrides_success_page(tmp_path):
    m = monitor(tmp_path)
    accepted(m)
    m.responses[SAVE_PATH][0].status = 500
    with pytest.raises(TerminationSubmissionError):
        m.receipt(success_visible=True)


def test_save_business_error_is_not_hidden_by_http_200(tmp_path):
    m = monitor(tmp_path)
    accepted(m)
    m.responses[SAVE_PATH][0].text.return_value = '{"success": false, "message": "保存失败"}'
    m.responses[SAVE_PATH][0].json.return_value = {"success": False, "message": "保存失败"}
    with pytest.raises(TerminationSubmissionError):
        m.receipt(success_visible=True)


def test_missing_response_or_duplicate_submission_never_passes(tmp_path):
    m = monitor(tmp_path)
    assert m.receipt(success_visible=True) is None
    accepted(m)
    m.responses[APPLY_PATH].append(m.responses[APPLY_PATH][0])
    with pytest.raises(TerminationSubmissionError):
        m.receipt(success_visible=True)


def test_submit_service_fills_submits_and_reopens_for_each_person(tmp_path):
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    automation = Mock()
    items = [EmploymentTerminationItem(IDENTITY, "6304"), EmploymentTerminationItem("320101199002021235", "6303")]
    with patch("ehrm.modules.employment_termination.service.EmploymentTerminationPage", return_value=automation):
        result = EmploymentTerminationService(settings, Mock()).submit_with_page(Mock(), items)
    assert automation.open.call_count == 2
    assert automation.submit_item.call_count == 2
    assert automation.close_submitted_form.call_count == 1
    assert result.submitted is True
    assert [r.status for r in result.results] == ["submitted", "submitted"]
    calls = [c[0] for c in automation.method_calls]
    assert calls == ["open", "fill_item", "submit_item", "close_submitted_form", "open", "fill_item", "submit_item"]


def test_unknown_stops_batch_and_preserves_row_results(tmp_path):
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    automation = Mock()
    automation.submit_item.side_effect = TerminationSubmissionError("响应超时，不得重复提交", outcome="unknown")
    items = [EmploymentTerminationItem(IDENTITY, "6304"), EmploymentTerminationItem("320101199002021235", "6303")]
    with patch("ehrm.modules.employment_termination.service.EmploymentTerminationPage", return_value=automation):
        result = EmploymentTerminationService(settings, Mock()).submit_with_page(Mock(), items)
    assert automation.submit_item.call_count == 1
    automation.close_submitted_form.assert_not_called()
    assert [r.status for r in result.results] == ["unknown", "not_processed"]
    assert result.submitted is False


def test_submit_service_can_skip_missing_employee_without_submitting_them(tmp_path):
    from ehrm.core.exceptions import EmployeeNotFoundError
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    automation = Mock()
    automation.fill_item.side_effect = [EmployeeNotFoundError("未查询到人员"), None]
    items = [EmploymentTerminationItem(IDENTITY, "6304"), EmploymentTerminationItem("320101199002021235", "6303")]
    with patch("ehrm.modules.employment_termination.service.EmploymentTerminationPage", return_value=automation):
        result = EmploymentTerminationService(settings, Mock(), progress_callback=Mock()).submit_with_page(Mock(), items)
    assert [r.status for r in result.results] == ["failed", "submitted"]
    assert automation.submit_item.call_count == 1
    automation.dismiss_query_feedback.assert_called_once()


def test_success_is_preserved_when_cleanup_fails(tmp_path):
    from playwright.sync_api import Error
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    automation = Mock()
    automation.close_submitted_form.side_effect = Error("窗口未关闭")
    items = [EmploymentTerminationItem(IDENTITY, "6304"), EmploymentTerminationItem("320101199002021235", "6303")]
    with patch("ehrm.modules.employment_termination.service.EmploymentTerminationPage", return_value=automation):
        result = EmploymentTerminationService(settings, Mock()).submit_with_page(Mock(), items)
    assert [r.status for r in result.results] == ["submitted", "not_processed"]
    assert result.submitted_count == 1
    assert automation.submit_item.call_count == 1


def test_workbook_distinguishes_submission_and_unknown_and_unprocessed(tmp_path):
    from openpyxl import Workbook, load_workbook
    from ehrm.modules.employment_termination.models import EmploymentTerminationItemResult
    from ehrm.modules.employment_termination.result_workbook import EmploymentTerminationResultWorkbookWriter
    source = tmp_path / "test.xlsx"
    book = Workbook()
    book.active.append(["身份证号", "退工原因"])
    for _ in range(3):
        book.active.append([IDENTITY, "6304"])
    book.save(source)
    book.close()
    rows = tuple(EmploymentTerminationItemResult(EmploymentTerminationItem(IDENTITY, "6304", index + 2), success, "code", "reason", status)
                 for index, (success, status) in enumerate([(True, "submitted"), (False, "unknown"), (False, "not_processed")]))
    target = EmploymentTerminationResultWorkbookWriter().write(source, tmp_path / "output", rows)
    result = load_workbook(target)
    try:
        assert [result.active.cell(row, 3).value for row in range(2, 5)] == ["提交成功", "结果待确认", "未执行"]
    finally:
        result.close()


def test_cli_submission_requires_explicit_flag_and_rejects_pause_before_fill():
    from ehrm.entrypoints.employment_termination_e2e_cli import _parser
    assert not _parser().parse_args(["--input", "test.xlsx"]).submit
    assert _parser().parse_args(["--input", "test.xlsx", "--submit"]).submit
    with pytest.raises(SystemExit):
        _parser().parse_args(["--input", "test.xlsx", "--submit", "--pause-after-fill"])


def run_submitter(tmp_path, *, success=True, dialog_error=False, native_error=False, stacked_error=False, late_error=False, inner_replaced=False):
    from dataclasses import replace
    from ehrm.browser.smart_wait import SmartWait
    from ehrm.modules.employment_termination.submission import TerminationSubmitter
    from ehrm.modules.jshrss_business_dialog import BusinessDialog, BusinessDialogKind

    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    settings = replace(settings, employment_termination=replace(settings.employment_termination, response_timeout_ms=1000))
    state = {"confirmed": False, "clicked": False, "clock": 0, "notice_closed": False}
    page, outer, frame = Mock(), Mock(), Mock()
    frame.parent_frame = outer
    listeners = {}
    page.on.side_effect = lambda name, listener: listeners.setdefault(name, listener)
    submit_button, yes_button = Mock(), Mock()
    yes_button.count.return_value = 1
    yes_button.is_visible.return_value = True
    no_button = Mock()
    no_button.count.return_value = 1
    no_button.is_visible.return_value = True
    declaration_locator = Mock()
    declaration_locator.get_by_role.side_effect = lambda role, name, exact: yes_button if name == "是" else no_button
    declaration = BusinessDialog(declaration_locator, None, "消息提示", settings.employment_termination.submit_declaration_text, BusinessDialogKind.INFORMATION)
    error = BusinessDialog(Mock(), None, "消息提示", "核心处理异常", BusinessDialogKind.ERROR)
    notice = BusinessDialog(Mock(), Mock(), "消息提示", "当前时间办理人员停保，发送税务缴费截止时间", BusinessDialogKind.INFORMATION)

    def role_locator(role, name, **kwargs):
        if name == settings.employment_termination.submit_button_text:
            return submit_button
        locator = Mock()
        locator.is_visible.side_effect = lambda: state["confirmed"] and success and (not inner_replaced or state["clock"] > 0)
        return locator

    outer.get_by_role.side_effect = role_locator
    def inner_role(role, name, **kwargs):
        if inner_replaced and state["confirmed"]:
            from playwright.sync_api import Error
            raise Error("Frame was detached")
        return role_locator(role, name, **kwargs)
    frame.get_by_role.side_effect = inner_role
    outer.get_by_text.return_value.is_visible.side_effect = lambda: state["confirmed"] and success
    frame.get_by_text.return_value.is_visible.side_effect = lambda: state["confirmed"] and success
    frame.locator.return_value.get_by_role.return_value.first.input_value.return_value = IDENTITY
    submit_button.click.side_effect = lambda **kwargs: state.update(clicked=True)

    def click_yes(**kwargs):
        state["confirmed"] = True
        # The registered listener observes all nested iframe requests.
        recorder = Mock(capture=listeners["response"])
        accepted(recorder)
        if native_error:
            native = Mock(message="服务器处理异常")
            listeners["dialog"](native)
            native.dismiss.assert_called_once()

    yes_button.click.side_effect = click_yes
    automation = Mock(page=page, settings=settings, contract=settings.employment_termination)
    automation.pacer.perform.side_effect = lambda action: action()
    automation.waiter = SmartWait(lambda ms: state.update(clock=state["clock"] + ms / 1000), clock=lambda: state["clock"])

    def dialog(scope):
        if scope is not frame or not state["clicked"]:
            return None
        if not state["confirmed"]:
            return declaration
        if stacked_error:
            return error if state["notice_closed"] else notice
        if late_error and state["clock"] > 0:
            return error
        return error if dialog_error else None

    with (patch("ehrm.modules.employment_termination.submission.topmost_business_dialog", side_effect=dialog),
          patch("ehrm.modules.employment_termination.submission.close_business_dialog", side_effect=lambda *a, **kw: state.update(notice_closed=True))):
        try:
            return TerminationSubmitter(automation, frame, EmploymentTerminationItem(IDENTITY, "6304")).run()
        finally:
            submit_button.click.assert_called_once()
            yes_button.click.assert_called_once()
            assert page.remove_listener.call_count == 2
            assert automation.pacer.perform.call_count == 2


def test_submitter_clicks_once_and_observes_network_and_dom(tmp_path):
    assert run_submitter(tmp_path).application_id == "application-1"


def test_submitter_waits_for_success_if_inner_form_is_replaced(tmp_path):
    assert run_submitter(tmp_path, inner_replaced=True).application_id == "application-1"


@pytest.mark.parametrize("kwargs", [{"success": False}, {"dialog_error": True}, {"native_error": True}, {"stacked_error": True}, {"late_error": True}])
def test_submitter_stops_without_reclicking_when_result_is_missing_or_error(tmp_path, kwargs):
    with pytest.raises(TerminationSubmissionError):
        run_submitter(tmp_path, **kwargs)


def test_real_error_message_must_not_be_ignored_even_when_ok_is_true(tmp_path):
    m = monitor(tmp_path)
    accepted(m)
    m.responses[APPLY_PATH][0].json.return_value["errmsg"] = "提交失败"
    with pytest.raises(TerminationSubmissionError):
        m.receipt(success_visible=True)


def test_shared_dialog_classifier_detects_submit_failure_under_generic_title():
    from ehrm.modules.jshrss_business_dialog import BusinessDialogKind, classify_business_dialog
    assert classify_business_dialog("消息提示", "提交失败，请查看日志") is BusinessDialogKind.ERROR


def test_submit_cli_routes_explicit_mode_and_exports_failure(tmp_path):
    from ehrm.core.auth_repository import AuthenticationRepository, SystemType
    from ehrm.entrypoints.employment_termination_e2e_cli import main
    from ehrm.modules.employment_termination.models import EmploymentTerminationPreparation, EmploymentTerminationItemResult
    from openpyxl import Workbook

    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    AuthenticationRepository(settings.auth_database_path).save_account(SystemType.JSHRSS, "test-unit", "test-password", secondary_account="test-mobile")
    path = tmp_path / "test.xlsx"
    workbook = Workbook()
    workbook.active.append(["身份证号", "退工原因"])
    workbook.active.append([IDENTITY, "6304"])
    workbook.save(path)
    workbook.close()
    item = EmploymentTerminationItem(IDENTITY, "6304", 2)
    result = EmploymentTerminationPreparation((item,), (EmploymentTerminationItemResult(item, False, "TERMINATION_RESULT_UNKNOWN", "超时", "unknown"),), "320100", "南京")
    service = Mock()
    service.submit_with_page.return_value = result
    browser = MagicMock()
    browser.__enter__.return_value = browser
    login = Mock()
    login.page = browser.page
    with (patch("ehrm.entrypoints.employment_termination_e2e_cli.load_settings", return_value=settings),
          patch("ehrm.entrypoints.employment_termination_e2e_cli.application_runtime_root", return_value=tmp_path),
          patch("ehrm.entrypoints.employment_termination_e2e_cli.BrowserManager", return_value=browser),
          patch("ehrm.entrypoints.employment_termination_e2e_cli.LoginService", return_value=login),
          patch("ehrm.entrypoints.employment_termination_e2e_cli.EmploymentTerminationService", return_value=service),
          patch("ehrm.entrypoints.employment_termination_e2e_cli._save_screenshot"),
          patch("builtins.input", return_value="") as prompt):
        assert main(["--input", str(path), "--submit", "--keep-browser-open"]) == 1
    service.prepare_with_page.assert_not_called()
    service.submit_with_page.assert_called_once()
    prompt.assert_called_once()
    assert list((tmp_path / "output/employment-termination-e2e").glob("*.xlsx"))
