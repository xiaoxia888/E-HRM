from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook

from ehrm.core.auth_repository import AuthenticationRepository, SystemType
from ehrm.core.error_catalog import ErrorCode
from ehrm.core.exceptions import EmployeeNotFoundError
from ehrm.core.settings import load_settings
from ehrm.entrypoints.employment_termination_e2e_cli import main
from ehrm.modules.employment_termination.models import (
    EmploymentTerminationItem,
    EmploymentTerminationItemResult,
    EmploymentTerminationPreparation,
)


def test_check_input_reads_saved_sqlite_account_without_prompt(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    AuthenticationRepository(settings.auth_database_path).save_account(
        SystemType.JSHRSS,
        "test-unit-code",
        "test-password",
        secondary_account="test-mobile",
    )
    workbook_path = tmp_path / "退保测试.xlsx"
    workbook = Workbook()
    workbook.active.append(["身份证号", "退工原因"])
    workbook.active.append(["320101199001011234", "6304"])
    workbook.save(workbook_path)
    workbook.close()

    with (
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.application_runtime_root",
            return_value=tmp_path,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.load_settings",
            return_value=settings,
        ),
    ):
        status = main(["--input", str(workbook_path), "--check-input"])

    assert status == 0


def test_default_system_remains_direct_when_government_account_is_database_default(
    tmp_path: Path,
    capsys: object,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    repository = AuthenticationRepository(settings.auth_database_path)
    repository.save_account(
        SystemType.JSHRSS,
        "direct-unit-code",
        "direct-password",
        secondary_account="direct-mobile",
        display_name="直登账号",
        profile={"login_method": "direct"},
    )
    repository.save_account(
        SystemType.JSHRSS,
        "government-user",
        "government-password",
        secondary_account="",
        display_name="政务账号",
        profile={"login_method": "government_sso"},
    )
    workbook_path = tmp_path / "退保测试.xlsx"
    workbook = Workbook()
    workbook.active.append(["身份证号", "退工原因"])
    workbook.active.append(["320101199001011234", "6304"])
    workbook.save(workbook_path)
    workbook.close()

    with (
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.application_runtime_root",
            return_value=tmp_path,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.load_settings",
            return_value=settings,
        ),
    ):
        status = main(["--input", str(workbook_path), "--check-input"])

    assert status == 0
    output = capsys.readouterr().out
    assert "直登账号" in output
    assert "方式=智慧人社直登" in output


def test_check_input_accepts_government_login_system(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    account = AuthenticationRepository(settings.auth_database_path).save_account(
        SystemType.JSHRSS,
        "government-user",
        "test-password",
        secondary_account="",
        display_name="政务服务平台账号",
        profile={"login_method": "government_sso"},
    )
    workbook_path = tmp_path / "退保测试.xlsx"
    workbook = Workbook()
    workbook.active.append(["身份证号", "退工原因"])
    workbook.active.append(["320101199001011234", "6304"])
    workbook.save(workbook_path)
    workbook.close()

    with (
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.application_runtime_root",
            return_value=tmp_path,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.load_settings",
            return_value=settings,
        ),
    ):
        status = main(
            [
                "--input",
                str(workbook_path),
                "--system",
                "江苏政务服务平台账号",
                "--check-input",
            ]
        )

    assert status == 0


def test_selected_government_account_controls_termination_login(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    account = AuthenticationRepository(settings.auth_database_path).save_account(
        SystemType.JSHRSS,
        "government-user",
        "test-password",
        secondary_account="",
        display_name="政务服务平台账号",
        profile={"login_method": "government_sso"},
    )
    workbook_path = tmp_path / "退保测试.xlsx"
    workbook = Workbook()
    workbook.active.append(["身份证号", "退工原因"])
    workbook.active.append(["320101199001011234", "6304"])
    workbook.save(workbook_path)
    workbook.close()

    class FakeBrowserManager:
        def __init__(self) -> None:
            self.page = Mock()

        def __enter__(self) -> "FakeBrowserManager":
            return self

        def __exit__(self, *_args: object) -> None:
            return None

    captured: dict[str, object] = {}

    class CapturingLoginService:
        def __init__(
            self,
            page: object,
            runtime_settings: object,
            **_kwargs: object,
        ) -> None:
            self.page = page
            captured["settings"] = runtime_settings

        def ensure_authenticated(self) -> None:
            raise ValueError("stop after checking selected login settings")

    with (
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.application_runtime_root",
            return_value=tmp_path,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.load_settings",
            return_value=settings,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.BrowserManager",
            return_value=FakeBrowserManager(),
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.LoginService",
            CapturingLoginService,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.configure_logging",
            return_value=Mock(),
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli._save_screenshot",
        ),
    ):
        status = main(
            [
                "--input",
                str(workbook_path),
                "--system",
                "江苏政务服务平台账号",
            ]
        )

    runtime_settings = captured["settings"]
    assert status == 1
    assert runtime_settings.rights_credentials.login_method == "government_sso"
    assert runtime_settings.rights_credentials.account_id == account.id
    assert runtime_settings.rights_credentials.credit_code == "government-user"
    assert runtime_settings.rights_credentials.mobile == ""
    assert runtime_settings.browser.user_data_dir.name == f"jshrss-{account.id}"


def test_failure_screenshot_is_saved_before_browser_closes(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    AuthenticationRepository(settings.auth_database_path).save_account(
        SystemType.JSHRSS,
        "test-unit-code",
        "test-password",
        secondary_account="test-mobile",
    )
    workbook_path = tmp_path / "退保测试.xlsx"
    workbook = Workbook()
    workbook.active.append(["身份证号", "退工原因"])
    workbook.active.append(["320101199001011234", "6304"])
    workbook.save(workbook_path)
    workbook.close()

    class FakeBrowserManager:
        def __init__(self) -> None:
            self.page = Mock()
            self.active = False

        def __enter__(self) -> "FakeBrowserManager":
            self.active = True
            return self

        def __exit__(self, *_args: object) -> None:
            self.active = False

    browser = FakeBrowserManager()

    class FakeLoginService:
        def __init__(self, page: object, *_args: object, **_kwargs: object) -> None:
            self.page = page

        def ensure_authenticated(self) -> None:
            return None

    service = Mock()
    service.prepare_with_page.side_effect = EmployeeNotFoundError("人员不存在")

    def assert_browser_is_active(*_args: object, **_kwargs: object) -> None:
        assert browser.active is True

    with (
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.application_runtime_root",
            return_value=tmp_path,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.load_settings",
            return_value=settings,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.BrowserManager",
            return_value=browser,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.LoginService",
            FakeLoginService,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.EmploymentTerminationService",
            return_value=service,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.configure_logging",
            return_value=Mock(),
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli._save_screenshot",
            side_effect=assert_browser_is_active,
        ) as save_screenshot,
    ):
        status = main(["--input", str(workbook_path)])

    assert status == 1
    save_screenshot.assert_called_once()
    assert browser.active is False


def test_row_level_not_found_result_exports_excel_and_returns_success(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    AuthenticationRepository(settings.auth_database_path).save_account(
        SystemType.JSHRSS,
        "test-unit-code",
        "test-password",
        secondary_account="test-mobile",
    )
    workbook_path = tmp_path / "退保测试.xlsx"
    workbook = Workbook()
    workbook.active.append(["身份证号", "退工原因"])
    workbook.active.append(["320101199001011234", "6304"])
    workbook.save(workbook_path)
    workbook.close()

    class FakeBrowserManager:
        def __init__(self) -> None:
            self.page = Mock()

        def __enter__(self) -> "FakeBrowserManager":
            return self

        def __exit__(self, *_args: object) -> None:
            return None

    class FakeLoginService:
        def __init__(self, page: object, *_args: object, **_kwargs: object) -> None:
            self.page = page

        def ensure_authenticated(self) -> None:
            return None

    item = EmploymentTerminationItem("320101199001011234", "6304", 2)
    item_result = EmploymentTerminationItemResult(
        item=item,
        success=False,
        code=ErrorCode.EMPLOYEE_NOT_FOUND,
        message="未能找到该人员信息",
    )
    preparation = EmploymentTerminationPreparation(
        items=(item,),
        results=(item_result,),
        city_code="320100",
        city_name="南京",
    )
    service = Mock()
    service.prepare_with_page.return_value = preparation
    writer = Mock()
    result_path = tmp_path / "output" / "退保录入结果.xlsx"
    writer.write.return_value = result_path

    with (
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.application_runtime_root",
            return_value=tmp_path,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.load_settings",
            return_value=settings,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.BrowserManager",
            return_value=FakeBrowserManager(),
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.LoginService",
            FakeLoginService,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.EmploymentTerminationService",
            return_value=service,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.EmploymentTerminationResultWorkbookWriter",
            return_value=writer,
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli.configure_logging",
            return_value=Mock(),
        ),
        patch(
            "ehrm.entrypoints.employment_termination_e2e_cli._save_screenshot",
        ),
    ):
        status = main(["--input", str(workbook_path)])

    assert status == 0
    writer.write.assert_called_once()
    assert writer.write.call_args.args[0] == workbook_path
    assert writer.write.call_args.args[2] == (item_result,)
