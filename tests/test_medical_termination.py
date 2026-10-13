from datetime import date
from itertools import count
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, Mock, call, patch

from openpyxl import Workbook, load_workbook
import pytest

from ehrm.browser.login import LoginService
from ehrm.core.auth_repository import AuthenticationRepository, SystemType
from ehrm.core.exceptions import EmployeeNotFoundError, ExcelValidationError, WebsiteStructureChangedError
from ehrm.core.preferences import UserPreferences
from ehrm.core.settings import BrowserInteractionPacingSettings, load_settings
from ehrm.entrypoints.medical_termination_e2e_cli import (
    _message, _select_account, _validate_login_url, _write_result, main,
)
from ehrm.modules.medical_termination.excel_loader import MedicalTerminationExcelLoader
from ehrm.modules.medical_termination.models import MedicalTerminationItem
from ehrm.modules.medical_termination.page import (
    MedicalSubmissionUncertainError, MedicalTerminationPage,
    is_medical_portal_url,
    verify_existing_medical_session,
)
from ehrm.web.settings_service import WebSettingsService


def _workbook(path: Path, rows: list[list[object]]) -> Path:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    workbook.close()
    return path


def _item() -> MedicalTerminationItem:
    return MedicalTerminationItem(
        "320101199001011234", "人员中断缴费", date(2026, 10, 12), False, 2
    )


def test_excel_loader_requires_exactly_one_valid_row(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "医保停保.xlsx", [
        ["身份证号", "停保原因", "停保日期", "停保当月缴费"],
        ["320101199001011234", "人员中断缴费", "2026-10-12", "否"],
    ])
    item = MedicalTerminationExcelLoader().load(path)
    assert item == _item()

    extra = _workbook(tmp_path / "两人.xlsx", [
        ["身份证号", "停保原因", "停保日期", "停保当月缴费"],
        ["320101199001011234", "人员中断缴费", "2026-10-12", "否"],
        ["320101199001011235", "人员中断缴费", "2026-10-12", "否"],
    ])
    with pytest.raises(ExcelValidationError, match="仅允许一名人员"):
        MedicalTerminationExcelLoader().load(extra)


def test_medical_reason_does_not_use_social_security_code(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "错误原因.xlsx", [
        ["身份证号", "停保原因", "停保日期", "停保当月缴费"],
        ["320101199001011234", "6304", "2026-10-12", "否"],
    ])
    with pytest.raises(ExcelValidationError, match="未在录制页面选项中"):
        MedicalTerminationExcelLoader().load(path)


def test_portal_match_ignores_token_values() -> None:
    assert is_medical_portal_url(
        "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    )
    assert not is_medical_portal_url("https://example.org/hsa-local-nj/web/hallUnit/")


def test_query_waits_until_loading_mask_has_cleared() -> None:
    page = SimpleNamespace(
        url="https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home",
        wait_for_timeout=Mock(),
    )
    automation = MedicalTerminationPage(page)  # type: ignore[arg-type]
    states = iter([True, True, False, False, False])
    automation._any_visible = Mock(side_effect=lambda _selector: next(states))  # type: ignore[method-assign]
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    automation._form_visible = Mock(return_value=True)  # type: ignore[method-assign]
    clock = count(step=0.2)

    with patch(
        "ehrm.modules.medical_termination.page.time.monotonic",
        side_effect=lambda: next(clock),
    ):
        automation._wait_for_query(_item(), [])

    assert page.wait_for_timeout.call_count >= 3
    automation._form_visible.assert_called_once()


def test_query_response_fallback_still_waits_for_stable_mask_free_state() -> None:
    page = SimpleNamespace(
        url="https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home",
        wait_for_timeout=Mock(),
    )
    automation = MedicalTerminationPage(page)  # type: ignore[arg-type]
    automation._any_visible = Mock(return_value=False)  # type: ignore[method-assign]
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    automation._form_visible = Mock(return_value=True)  # type: ignore[method-assign]
    clock = count(step=0.2)
    response = SimpleNamespace(status=200)

    with patch(
        "ehrm.modules.medical_termination.page.time.monotonic",
        side_effect=lambda: next(clock),
    ):
        automation._wait_for_query(_item(), [response])  # type: ignore[list-item]

    assert page.wait_for_timeout.call_count >= 1


def test_business_error_is_not_treated_as_success() -> None:
    error = SimpleNamespace(is_visible=lambda: True, inner_text=lambda: "该人员不是本单位人员!")
    empty = SimpleNamespace(count=lambda: 0)
    page = SimpleNamespace(
        url="https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home",
        locator=lambda selector: SimpleNamespace(
            count=lambda: 1,
            nth=lambda _index: error,
        ) if selector == ".el-message-box__message" else empty,
    )
    automation = MedicalTerminationPage(page)  # type: ignore[arg-type]

    with pytest.raises(EmployeeNotFoundError, match="不是本单位人员"):
        automation._raise_feedback(_item())


def test_reason_waits_for_recorded_visible_listitem_before_click() -> None:
    input_field = SimpleNamespace(input_value=lambda: "单位调动")
    control = SimpleNamespace(
        is_visible=lambda: True,
        click=Mock(),
        locator=lambda _selector: SimpleNamespace(first=input_field),
    )
    controls = SimpleNamespace(count=lambda: 1, nth=lambda _index: control)
    option = SimpleNamespace(wait_for=Mock(), click=Mock())
    options = SimpleNamespace(first=option, count=lambda: 1)
    listitems = SimpleNamespace(filter=Mock(return_value=options))
    page = SimpleNamespace(
        url="https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home",
        locator=lambda _selector: controls,
        get_by_role=Mock(return_value=listitems),
        wait_for_timeout=Mock(),
    )
    automation = MedicalTerminationPage(  # type: ignore[arg-type]
        page, pacing=BrowserInteractionPacingSettings(True, 100, 100),
    )

    automation._select_option(".el-form-item.is-required .el-select", "单位调动", "停保原因")

    page.get_by_role.assert_called_once_with("listitem")
    option.wait_for.assert_called_once_with(state="visible", timeout=30_000)
    option.click.assert_called_once()
    control.click.assert_called_once()
    assert page.wait_for_timeout.call_args_list == [call(100), call(100)]


def test_submit_clicks_once_only_after_preflight_and_new_success_toast() -> None:
    page = MagicMock()
    page.url = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    page.get_by_role.return_value.first.input_value.side_effect = [
        "320101199001011234", "2026-10-12",
    ]
    page.get_by_role.return_value.count.return_value = 1
    page.get_by_role.return_value.nth.return_value.is_disabled.return_value = False
    page.locator.return_value.count.return_value = 1
    page.locator.return_value.nth.return_value.locator.return_value.first.input_value.side_effect = [
        "人员中断缴费", "否",
    ]
    automation = MedicalTerminationPage(page)
    automation._form_visible = Mock(return_value=True)  # type: ignore[method-assign]
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    automation._visible = Mock(return_value=True)  # type: ignore[method-assign]
    automation._visible_messages = Mock(side_effect=[[], ["提交成功！"]])  # type: ignore[method-assign]
    events = []

    result = automation.submit_item(_item(), before_submit=lambda: events.append("durable-submitting"))

    assert "平台显示" in result
    assert events == ["durable-submitting"]
    page.get_by_role.return_value.nth.return_value.click.assert_called_once()


def test_submit_error_after_click_is_uncertain_and_never_clicks_twice() -> None:
    page = MagicMock()
    page.url = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    page.get_by_role.return_value.first.input_value.side_effect = [
        "320101199001011234", "2026-10-12",
    ]
    page.get_by_role.return_value.count.return_value = 1
    page.get_by_role.return_value.nth.return_value.is_disabled.return_value = False
    page.locator.return_value.count.return_value = 1
    page.locator.return_value.nth.return_value.locator.return_value.first.input_value.side_effect = [
        "人员中断缴费", "否",
    ]
    automation = MedicalTerminationPage(page)
    automation._form_visible = Mock(return_value=True)  # type: ignore[method-assign]
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    automation._visible = Mock(return_value=True)  # type: ignore[method-assign]
    automation._visible_messages = Mock(return_value=["提交失败，请人工核对"])  # type: ignore[method-assign]

    with pytest.raises(MedicalSubmissionUncertainError, match="请勿重复提交"):
        automation.submit_item(_item())
    page.get_by_role.return_value.nth.return_value.click.assert_called_once()


def test_submit_preflight_mismatch_never_crosses_mutation_boundary() -> None:
    page = MagicMock()
    page.url = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    page.get_by_role.return_value.first.input_value.side_effect = [
        "320101199001011234", "2026-10-11",
    ]
    automation = MedicalTerminationPage(page)
    automation._form_visible = Mock(return_value=True)  # type: ignore[method-assign]
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    before = Mock()

    with pytest.raises(WebsiteStructureChangedError, match="日期与申请数据不一致"):
        automation.submit_item(_item(), before_submit=before)
    before.assert_not_called()
    page.get_by_role.return_value.nth.return_value.click.assert_not_called()


def test_submit_does_not_click_if_portal_identity_differs_from_application() -> None:
    page = MagicMock()
    page.url = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    page.get_by_role.return_value.first.input_value.return_value = "320101199002021235"
    automation = MedicalTerminationPage(page)
    automation._form_visible = Mock(return_value=True)  # type: ignore[method-assign]
    automation._raise_feedback = Mock()  # type: ignore[method-assign]
    before = Mock()

    with pytest.raises(WebsiteStructureChangedError, match="证件号码与本次人员不一致"):
        automation.submit_item(_item(), before_submit=before)
    before.assert_not_called()
    page.get_by_role.return_value.nth.return_value.click.assert_not_called()


def test_next_medical_person_requires_verified_session_and_empty_query_form() -> None:
    page = MagicMock()
    page.url = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    automation = MedicalTerminationPage(page)
    automation.open = Mock()  # type: ignore[method-assign]
    with patch("ehrm.modules.medical_termination.page.verify_existing_medical_session", return_value=True) as verify:
        page.get_by_role.return_value.first.input_value.return_value = ""
        automation.prepare_next()
        verify.assert_called_once_with(page, timeout_ms=30_000)
        automation.open.assert_called_once()

        page.get_by_role.return_value.first.input_value.return_value = _item().identity_number
        with pytest.raises(WebsiteStructureChangedError, match="仍保留上一人"):
            automation.prepare_next()

    with patch("ehrm.modules.medical_termination.page.verify_existing_medical_session", return_value=False):
        with pytest.raises(WebsiteStructureChangedError, match="会话未通过验证"):
            automation.prepare_next()


def test_check_input_uses_saved_government_account_without_browser(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    AuthenticationRepository(settings.auth_database_path).save_account(
        SystemType.MEDICAL,
        "government-user",
        "government-password",
        secondary_account="",
        profile={"login_method": "government_sso"},
    )
    path = _workbook(tmp_path / "医保停保.xlsx", [
        ["身份证号", "停保原因", "停保日期", "停保当月缴费"],
        ["320101199001011234", "人员中断缴费", "2026-10-12", "否"],
    ])
    with (
        patch(
            "ehrm.entrypoints.medical_termination_e2e_cli.application_runtime_root",
            return_value=tmp_path,
        ),
        patch(
            "ehrm.entrypoints.medical_termination_e2e_cli.load_settings",
            return_value=settings,
        ),
        patch("ehrm.entrypoints.medical_termination_e2e_cli.BrowserManager") as browser,
    ):
        status = main(["--input", str(path), "--check-input"])

    assert status == 0
    browser.assert_not_called()


def test_termination_cli_passes_saved_execution_speed_to_business_page(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    AuthenticationRepository(settings.auth_database_path).save_account(
        SystemType.MEDICAL, "government-user", "government-password",
        profile={"login_method": "government_sso"},
    )
    WebSettingsService(settings, tmp_path).preferences_store.save(
        UserPreferences(execution_speed="fast")
    )
    source = _workbook(tmp_path / "停保.xlsx", [
        ["身份证号", "停保原因", "停保日期", "停保当月缴费"],
        ["320101199001011234", "人员中断缴费", "2026-10-12", "否"],
    ])
    portal = SimpleNamespace(url="https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home")
    browser = MagicMock()
    browser.return_value.__enter__.return_value.page = portal
    with (
        patch("ehrm.entrypoints.medical_termination_e2e_cli.application_runtime_root", return_value=tmp_path),
        patch("ehrm.entrypoints.medical_termination_e2e_cli.load_settings", return_value=settings),
        patch("ehrm.entrypoints.medical_termination_e2e_cli.BrowserManager", browser),
        patch("ehrm.entrypoints.medical_termination_e2e_cli.LoginService") as login,
        patch("ehrm.entrypoints.medical_termination_e2e_cli.MedicalTerminationPage") as business_page,
        patch("ehrm.entrypoints.medical_termination_e2e_cli._save_screenshot"),
        patch("ehrm.entrypoints.medical_termination_e2e_cli._write_result", return_value=tmp_path / "result.xlsx"),
    ):
        login.return_value.ensure_external_government_authenticated.return_value = portal
        assert main(["--input", str(source)]) == 0

    business_page.return_value.fill_item.assert_called_once()
    configured = business_page.call_args.kwargs["pacing"]
    assert configured.enabled is True
    assert (configured.min_delay_ms, configured.max_delay_ms) == (500, 1000)


def test_medical_default_and_explicit_account_selection(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    repository = AuthenticationRepository(settings.auth_database_path)
    first = repository.save_account(
        SystemType.MEDICAL, "nanhua", "password-one",
        profile={"login_method": "government_sso"},
    )
    second = repository.save_account(
        SystemType.MEDICAL, "bahua", "password-two",
        profile={"login_method": "government_sso"},
    )
    assert _select_account(repository, None).id == second.id
    repository.set_default_account(SystemType.MEDICAL, first.id)
    assert _select_account(repository, None).id == first.id
    assert _select_account(repository, first.id).account == "nanhua"
    assert _select_account(repository, second.id).account == "bahua"


def test_social_security_government_account_is_not_a_medical_account(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    repository = AuthenticationRepository(settings.auth_database_path)
    repository.save_account(
        SystemType.JSHRSS, "social-user", "social-password",
        profile={"login_method": "government_sso"},
    )
    with pytest.raises(ValueError, match="医保政务服务平台账号"):
        _select_account(repository, None)


def test_medical_login_url_rejects_token_page_and_other_hosts() -> None:
    url = "https://www.jszwfw.gov.cn/jsjis/front/login.do?uuid=Q5mxXFYhCnUh"
    assert _validate_login_url(url) == url
    with pytest.raises(ValueError, match="官方 HTTPS"):
        _validate_login_url(
            "https://wb.njybjyybz.org.cn/hsa-local-nj/web/hallUnit/?accessToken=x"
        )


def test_medical_login_uses_direct_entry_and_not_jshrss_target(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    portal = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    page = SimpleNamespace(url="about:blank", on=Mock())
    page.goto = Mock(side_effect=lambda url, **_kw: setattr(page, "url", url))
    service = LoginService(page, settings)  # type: ignore[arg-type]

    def submit(*_args: object, **_kwargs: object) -> None:
        page.url = portal

    service._submit_government_login = Mock(side_effect=submit)  # type: ignore[method-assign]
    service._wait_for_external_entry_resolution = Mock(  # type: ignore[method-assign]
        return_value=(None, page)
    )
    entry = "https://www.jszwfw.gov.cn/jsjis/front/login.do?uuid=Q5mxXFYhCnUh"
    actual = service.ensure_external_government_authenticated(
        entry_url=entry,
        username="government-user",
        password="government-password",
        target_matches=is_medical_portal_url,
    )
    assert actual is page
    page.goto.assert_called_once_with(entry, wait_until="domcontentloaded")
    service._submit_government_login.assert_called_once()


def test_medical_login_accepts_direct_portal_redirect_without_form(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    portal = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    page = SimpleNamespace(url="about:blank", on=Mock())
    page.goto = Mock(side_effect=lambda _url, **_kw: setattr(page, "url", portal))
    service = LoginService(page, settings)  # type: ignore[arg-type]
    service._submit_government_login = Mock()  # type: ignore[method-assign]
    actual = service.ensure_external_government_authenticated(
        entry_url="https://www.jszwfw.gov.cn/jsjis/front/login.do?uuid=entry",
        username="government-user",
        password="government-password",
        target_matches=is_medical_portal_url,
    )
    assert actual is page
    service._submit_government_login.assert_not_called()


def test_medical_login_reuses_verified_open_portal_without_visiting_sso(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    portal = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    page = SimpleNamespace(url=portal, on=Mock(), goto=Mock())
    service = LoginService(page, settings)  # type: ignore[arg-type]
    service._submit_government_login = Mock()  # type: ignore[method-assign]
    verify = Mock(return_value=True)

    actual = service.ensure_external_government_authenticated(
        entry_url="https://www.jszwfw.gov.cn/jsjis/front/login.do?uuid=entry",
        username="government-user",
        password="government-password",
        target_matches=is_medical_portal_url,
        existing_session_ready=verify,
    )

    assert actual is page
    verify.assert_called_once_with(page)
    page.goto.assert_not_called()
    service._submit_government_login.assert_not_called()


def test_medical_login_falls_back_to_sso_when_open_portal_is_stale(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    portal = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    entry = "https://www.jszwfw.gov.cn/jsjis/front/login.do?uuid=entry"
    page = SimpleNamespace(url=portal, on=Mock())
    page.goto = Mock(side_effect=lambda url, **_kw: setattr(page, "url", url))
    service = LoginService(page, settings)  # type: ignore[arg-type]
    service._wait_for_external_entry_resolution = Mock(  # type: ignore[method-assign]
        return_value=(None, page)
    )
    service._submit_government_login = Mock(  # type: ignore[method-assign]
        side_effect=lambda *_args, **_kwargs: setattr(page, "url", portal)
    )

    actual = service.ensure_external_government_authenticated(
        entry_url=entry,
        username="government-user",
        password="government-password",
        target_matches=is_medical_portal_url,
        existing_session_ready=Mock(return_value=False),
    )

    assert actual is page
    page.goto.assert_called_once_with(entry, wait_until="domcontentloaded")
    service._submit_government_login.assert_called_once()


def test_existing_medical_session_requires_successful_menu_reload() -> None:
    portal = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    response = SimpleNamespace(
        url="https://wb.njybjyybz.org.cn:9091/hsa-pss-cw-local-nj/pss/web/roles/selectCwMenu",
        request=SimpleNamespace(method="POST"),
        status=200,
        json=Mock(return_value={"code": 0, "data": {"success": True}}),
    )
    pending = MagicMock()
    pending.__enter__.return_value = SimpleNamespace(value=response)
    menu = SimpleNamespace(first=SimpleNamespace(wait_for=Mock(), is_visible=lambda: True))
    page = SimpleNamespace(
        url=portal,
        expect_response=Mock(return_value=pending),
        reload=Mock(),
        get_by_text=Mock(return_value=menu),
    )

    assert verify_existing_medical_session(page)  # type: ignore[arg-type]
    predicate = page.expect_response.call_args.args[0]
    assert predicate(response)
    page.reload.assert_called_once_with(wait_until="domcontentloaded", timeout=10_000)

    response.json.return_value = {"code": -1, "data": {"success": False}}
    assert not verify_existing_medical_session(page)  # type: ignore[arg-type]


def test_medical_portal_ready_requires_loaded_business_menu() -> None:
    from ehrm.modules.medical_termination.page import is_medical_portal_ready

    page = SimpleNamespace(
        url="https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home",
        get_by_text=Mock(return_value=SimpleNamespace(
            first=SimpleNamespace(is_visible=lambda: False)
        )),
    )
    assert not is_medical_portal_ready(page)  # type: ignore[arg-type]
    page.get_by_text.return_value.first.is_visible = lambda: True
    assert is_medical_portal_ready(page)  # type: ignore[arg-type]


def test_existing_sso_session_can_redirect_after_entry_load(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    portal = "https://wb.njybjyybz.org.cn:9091/hsa-local-nj/web/hallUnit/#/Home"
    landing = "https://www.jszwfw.gov.cn/jsjis/front/login.do?uuid=entry"
    visible = SimpleNamespace(first=SimpleNamespace(is_visible=lambda: True))
    page = SimpleNamespace(url=landing, on=Mock(), locator=lambda _selector: visible)
    waits = 0

    def advance(_milliseconds: int) -> None:
        nonlocal waits
        waits += 1
        if waits == 4:
            page.url = portal

    page.wait_for_timeout = Mock(side_effect=advance)
    service = LoginService(page, settings)  # type: ignore[arg-type]

    target, login_form = service._wait_for_external_entry_resolution(
        is_medical_portal_url, None
    )

    assert target is page
    assert login_form is None
    assert waits == 4


def test_expired_sso_session_stays_on_login_form(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    landing = "https://www.jszwfw.gov.cn/jsjis/front/login.do?uuid=entry"
    visible = SimpleNamespace(first=SimpleNamespace(is_visible=lambda: True))
    page = SimpleNamespace(
        url=landing, on=Mock(), locator=lambda _selector: visible,
        wait_for_timeout=Mock(),
    )
    service = LoginService(page, settings)  # type: ignore[arg-type]

    target, login_form = service._wait_for_external_entry_resolution(
        is_medical_portal_url, None
    )

    assert target is None
    assert login_form is page
    assert page.wait_for_timeout.call_count >= 14


def test_result_workbook_marks_prepared_not_submitted(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", [
        ["身份证号", "停保原因", "停保日期", "停保当月缴费"],
        ["320101199001011234", "人员中断缴费", "2026-10-12", "否"],
    ])
    result = _write_result(source, tmp_path / "results", "录入成功（未提交）", "")
    workbook = load_workbook(result, read_only=True)
    try:
        assert workbook.active.cell(2, 5).value == "录入成功（未提交）"
    finally:
        workbook.close()


def test_error_messages_do_not_expose_medical_portal_tokens() -> None:
    message = _message(ValueError(
        "url=#/Home?accessToken=secret-one&refreshToken=secret-two"
    ))
    assert "secret-one" not in message
    assert "secret-two" not in message
