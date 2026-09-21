from dataclasses import replace
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook
import pytest

from ehrm.core.auth_repository import AuthenticationRepository, SystemType
from ehrm.core.exceptions import ConfigurationError, EmployeeNotFoundError, WebsiteStructureChangedError
from ehrm.core.exceptions import ExcelValidationError, QueryValidationError
from ehrm.core.settings import load_settings
from ehrm.entrypoints.person_information_collection_e2e_cli import load_test_excel, main
from ehrm.modules.person_information_collection.models import (
    PersonInformationRowError,
    PersonInformationItem,
    address_level_for_household,
    normalize_items,
)
from ehrm.modules.person_information_collection.page import PersonInformationPage
from ehrm.modules.person_information_collection.service import PersonInformationService


def _item(**changes: object) -> PersonInformationItem:
    values = {
        "identity_number": "320101199001011234",
        "name": "测试人员",
        "mobile": "13800138000",
        "nation": "汉族",
        "household_type": "11",
        "address": ("江苏省", "南京市", "玄武区", "玄武门街道", "社区甲"),
        "source_index": 2,
    }
    values.update(changes)
    return PersonInformationItem(**values)


def test_domain_normalizes_identity_and_address() -> None:
    item = _item(
        identity_number=" 32010119900101123x ",
        address=(" 江苏省 ", " 南京市 ", "玄武区", "玄武门街道", "社区甲"),
    ).normalized()
    assert item.identity_number == "32010119900101123X"
    assert item.address[0] == "江苏省"


@pytest.mark.parametrize(
    "changes",
    [
        {"identity_number": "123"},
        {"name": ""},
        {"mobile": "123"},
        {"nation": ""},
        {"household_type": ""},
        {"address": ("江苏省", "", "玄武区")},
    ],
)
def test_domain_rejects_invalid_or_incomplete_form_data(changes: dict[str, object]) -> None:
    with pytest.raises(QueryValidationError):
        _item(**changes).normalized()


def test_domain_rejects_duplicate_people_before_browser() -> None:
    with pytest.raises(QueryValidationError, match="重复"):
        normalize_items([_item(), _item(source_index=3)])


@pytest.mark.parametrize(
    ("household", "expected"),
    [
        ("本省城镇", 5),
        ("本省农村", 5),
        ("外省城镇", 3),
        ("外省农村", 3),
        ("香港特别行政区", 0),
        ("澳门特别行政区", 0),
        ("台湾", 0),
        ("外国人", 0),
    ],
)
def test_household_label_decides_required_address_depth(
    household: str, expected: int
) -> None:
    assert address_level_for_household(household) == expected


def test_domain_allows_empty_three_or_five_level_address() -> None:
    assert _item(address=()).normalized().address == ()
    assert _item(address=("安徽省", "合肥市", "蜀山区", "", "")).normalized().address == (
        "安徽省", "合肥市", "蜀山区"
    )
    assert len(_item().normalized().address) == 5


def test_test_entrypoint_reads_excel_to_object_array(tmp_path: Path) -> None:
    path = tmp_path / "采集测试.xlsx"
    workbook = Workbook()
    workbook.active.append(
        ["身份证号", "姓名", "手机号", "民族", "户籍性质", "省", "市", "区县", "街道", "社区村"]
    )
    workbook.active.append(
        ["320101199001011234", "测试人员", "13800138000", "汉族", "11", "江苏省", "南京市", "玄武区", "玄武门街道", "社区甲"]
    )
    workbook.save(path)
    workbook.close()

    assert load_test_excel(path) == [_item()]


def test_excel_accepts_special_region_without_address_and_external_three_levels(
    tmp_path: Path,
) -> None:
    path = tmp_path / "采集测试.xlsx"
    workbook = Workbook()
    workbook.active.append(
        ["身份证号", "姓名", "手机号", "民族", "户籍性质", "省", "市", "区县", "街道", "社区村"]
    )
    workbook.active.append(
        ["320101199001011234", "港澳台示例", "13800138000", "汉族", "香港特别行政区", None, None, None, None, None]
    )
    workbook.active.append(
        ["320101199002021235", "外省示例", "13800138001", "土家族", "外省城镇", "安徽省", "合肥市", "蜀山区", None, None]
    )
    workbook.save(path)
    workbook.close()

    items = load_test_excel(path)
    assert items[0].address == ()
    assert items[1].address == ("安徽省", "合肥市", "蜀山区")


def test_test_entrypoint_rejects_missing_column(tmp_path: Path) -> None:
    path = tmp_path / "采集测试.xlsx"
    workbook = Workbook()
    workbook.active.append(["身份证号", "姓名"])
    workbook.save(path)
    workbook.close()
    with pytest.raises(ExcelValidationError, match="缺少必要列"):
        load_test_excel(path)


def test_page_uses_shared_hall_navigation_before_waiting_for_form(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    page.url = "https://rs.jshrss.jiangsu.gov.cn/web/login"
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    automation.navigator.open_menu = Mock()
    automation._wait_for_business_frame = Mock(return_value=Mock())
    automation._wait_for_form_ready = Mock()

    automation.open()

    automation.navigator.open_menu.assert_called_once_with("人员基础信息采集")


def test_collection_rejects_non_nanjing_configuration_before_opening_site(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    settings = replace(
        settings,
        employment_termination=replace(
            settings.employment_termination,
            city_code="320200",
            city_name="无锡",
        ),
    )
    page = Mock()
    page.url = "https://rs.jshrss.jiangsu.gov.cn/web/login"
    page.wait_for_timeout = lambda _milliseconds: None
    with pytest.raises(ConfigurationError, match="南京"):
        PersonInformationPage(page, settings).open()
    page.goto.assert_not_called()


def test_first_visible_returns_first_visible_candidate(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock(url="https://rs.jshrss.jiangsu.gov.cn/index/")
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    hidden = Mock()
    hidden.count.return_value = 1
    hidden.is_visible.return_value = False
    visible = Mock()
    visible.count.return_value = 1
    visible.is_visible.return_value = True
    candidates = Mock()
    candidates.count.return_value = 2
    candidates.nth.side_effect = [hidden, visible]

    assert automation._first_visible(candidates) is visible


def test_excel_rejects_numeric_identity_to_prevent_rounded_value(tmp_path: Path) -> None:
    path = tmp_path / "采集测试.xlsx"
    workbook = Workbook()
    workbook.active.append(
        ["身份证号", "姓名", "手机号", "民族", "户籍性质", "省", "市", "区县", "街道", "社区村"]
    )
    workbook.active.append(
        [320101199001011234, "测试人员", "13800138000", "汉族", "11", "江苏省", "南京市", "玄武区", "玄武门街道", "社区甲"]
    )
    workbook.save(path)
    workbook.close()
    with pytest.raises(ExcelValidationError) as error:
        load_test_excel(path)
    assert "文本格式" in error.value.details


def test_address_confirmation_is_scoped_to_picker_component(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    page.url = "https://rs.jshrss.jiangsu.gov.cn/index/"
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    automation.pacer.perform = Mock(side_effect=lambda action: action())
    automation.waiter.first = Mock()
    automation._visible = Mock(return_value=True)
    picker = Mock()
    component = picker.locator.return_value
    confirms = component.locator.return_value
    confirms.count.return_value = 1
    button = confirms.nth.return_value
    address_input = Mock()

    automation._confirm_address(
        picker,
        address_input,
        ("江苏省", "南京市", "玄武区", "玄武门街道", "社区甲"),
    )

    picker.locator.assert_called_once_with(
        "xpath=ancestor::*[contains(concat(' ', normalize-space(@class), ' '), ' girder-select-address ')][1]"
    )
    component.locator.assert_called_once_with("button.address-submit")
    button.click.assert_called_once()

    condition = automation.waiter.first.call_args.args[0][0]
    address_input.input_value.return_value = "江苏省南京市玄武区玄武门街道社区甲"
    assert condition.probe() is True


@pytest.mark.parametrize("requested", ["土家族", "15", "string:15"])
def test_nation_selects_live_option_by_name_or_code(
    tmp_path: Path,
    requested: str,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock(url="https://rs.jshrss.jiangsu.gov.cn/index/")
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    automation.pacer.perform = Mock(side_effect=lambda action: action())
    select = Mock()
    options = select.locator.return_value
    options.count.return_value = 2
    placeholder = options.nth.return_value
    nation = Mock()
    options.nth.side_effect = [placeholder, nation]
    placeholder.get_attribute.return_value = ""
    placeholder.inner_text.return_value = "-- 请选择 --"
    nation.get_attribute.return_value = "string:15"
    nation.inner_text.return_value = "土家族"
    select.input_value.return_value = "string:15"

    selected_label = automation._select_option(
        select,
        requested,
        field_name="民族",
        reason_code="NATION_NOT_FOUND",
    )

    assert selected_label == "土家族"
    select.select_option.assert_called_once_with(value="string:15")


def test_address_confirmation_does_not_require_picker_dom_to_be_hidden(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock(url="https://rs.jshrss.jiangsu.gov.cn/index/")
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    automation.pacer.perform = Mock(side_effect=lambda action: action())
    automation._visible = Mock(return_value=True)
    picker = Mock()
    confirms = picker.locator.return_value.locator.return_value
    confirms.count.return_value = 1
    address_input = Mock()
    address_input.input_value.return_value = "江苏省南京市玄武区玄武门街道社区甲"

    automation._confirm_address(
        picker,
        address_input,
        ("江苏省", "南京市", "玄武区", "玄武门街道", "社区甲"),
    )

    confirms.nth.return_value.click.assert_called_once()


def test_address_option_waits_for_async_level_to_appear(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    page.url = "https://rs.jshrss.jiangsu.gov.cn/index/"
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    automation._visible = Mock(return_value=True)
    option = Mock()
    automation._last_visible = Mock(side_effect=[None, option])
    automation._visible_address_columns = Mock(return_value=[Mock()])
    picker = Mock()

    assert automation._wait_for_address_option(
        picker, "江苏省", level=0, selected=()
    ) is option
    assert automation._last_visible.call_count == 2


def test_address_confirmation_refuses_unassociated_global_button(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    page.url = "https://rs.jshrss.jiangsu.gov.cn/index/"
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    automation._visible = Mock(return_value=True)
    picker = Mock()
    picker.locator.return_value.locator.return_value.count.return_value = 0

    with pytest.raises(WebsiteStructureChangedError, match="户籍地址控件内"):
        automation._confirm_address(
            picker,
            Mock(),
            ("江苏省", "南京市", "玄武区", "玄武门街道", "社区甲"),
        )


def test_missing_address_option_returns_row_problem_with_level_and_path(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock(url="https://rs.jshrss.jiangsu.gov.cn/index/")
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    automation._visible = Mock(return_value=True)
    automation._last_visible = Mock(return_value=None)
    automation._visible_address_columns = Mock(
        return_value=[Mock(), Mock(), Mock(), Mock(), Mock()]
    )
    automation.settings = replace(
        settings,
        browser=replace(settings.browser, action_timeout_ms=5),
    )
    with pytest.raises(PersonInformationRowError, match="第 5 级.*社区A"):
        automation._wait_for_address_option(
            Mock(), "社区A", level=4, selected=("江苏省", "南京市", "玄武区", "街道A")
        )


def test_address_policy_rejects_wrong_depth_before_opening_picker(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock(url="https://rs.jshrss.jiangsu.gov.cn/index/")
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    with pytest.raises(PersonInformationRowError, match="外省.*3 级"):
        automation._validate_address_depth(
            "外省城镇", ("安徽省", "合肥市", "蜀山区", "街道A")
        )
    with pytest.raises(PersonInformationRowError, match="无需填写"):
        automation._validate_address_depth("香港特别行政区", ("香港",))


def test_service_returns_business_popup_and_address_problem_then_continues(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    page.context.new_page.return_value = Mock()
    with patch(
        "ehrm.modules.person_information_collection.service.PersonInformationPage"
    ) as automation_class:
        automation_class.return_value.fill_item.side_effect = [
            PersonInformationRowError("页面提示：证件号码已存在", reason_code="BUSINESS_MESSAGE"),
            PersonInformationRowError("第 5 级没有找到社区A", reason_code="ADDRESS_NOT_FOUND"),
            None,
        ]
        result = PersonInformationService(settings, Mock()).prepare_with_page(
            page,
            [
                _item(),
                _item(identity_number="320101199002021235", source_index=3),
                _item(identity_number="320101199003031236", source_index=4),
            ],
        )
    assert [row.success for row in result.results] == [False, False, True]
    assert [row.code for row in result.results] == [
        "BUSINESS_MESSAGE", "ADDRESS_NOT_FOUND", "SUCCESS"
    ]
    assert "证件号码已存在" in result.results[0].message


def test_visible_business_popup_is_closed_and_returned_as_row_problem(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock(url="https://rs.jshrss.jiangsu.gov.cn/index/")
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    frame = Mock()
    dialogs = frame.locator.return_value
    dialogs.count.return_value = 1
    dialog = dialogs.nth.return_value
    body = dialog.locator.return_value
    body.inner_text.return_value = "证件号码已在系统中存在"
    automation._visible = Mock(return_value=True)
    automation._close_feedback_dialog = Mock()

    with pytest.raises(PersonInformationRowError) as error:
        automation._raise_on_feedback(frame)

    assert error.value.reason_code == "BUSINESS_MESSAGE"
    assert "证件号码已在系统中存在" in error.value.message
    automation._close_feedback_dialog.assert_called_once_with(dialog)


def test_address_selection_stops_at_live_leaf_and_confirms(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock(url="https://rs.jshrss.jiangsu.gov.cn/index/")
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    automation.pacer.perform = Mock(side_effect=lambda action: action())
    automation._wait_for_address_option = Mock(side_effect=[Mock(), Mock(), Mock()])
    automation._wait_for_address_step = Mock(side_effect=[True, True, False])
    automation._confirm_address = Mock()
    frame = Mock()
    frame.locator.return_value.wait_for.return_value = None

    automation._select_address(frame, ("安徽省", "合肥市", "蜀山区"))

    assert automation._wait_for_address_step.call_count == 3
    automation._confirm_address.assert_called_once_with(
        frame.locator.return_value,
        frame.get_by_role.return_value.nth.return_value,
        ("安徽省", "合肥市", "蜀山区"),
    )


def test_address_selection_reports_when_platform_ends_before_imported_path(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock(url="https://rs.jshrss.jiangsu.gov.cn/index/")
    page.wait_for_timeout = lambda _milliseconds: None
    automation = PersonInformationPage(page, settings)
    automation.pacer.perform = Mock(side_effect=lambda action: action())
    automation._wait_for_address_option = Mock(return_value=Mock())
    automation._wait_for_address_step = Mock(return_value=False)
    frame = Mock()

    with pytest.raises(PersonInformationRowError) as error:
        automation._select_address(
            frame,
            ("江苏省", "南京市", "玄武区", "玄武门街道", "社区A"),
        )

    assert error.value.reason_code == "ADDRESS_LEVEL_ENDED"
    assert "南京市" in error.value.message


def test_service_accepts_objects_and_never_submits(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    page.wait_for_timeout = lambda _milliseconds: None
    logger = Mock()
    with patch(
        "ehrm.modules.person_information_collection.service.PersonInformationPage"
    ) as automation_class:
        preparation = PersonInformationService(settings, logger).prepare_with_page(
            page, [_item()]
        )
    automation_class.return_value.open.assert_called_once()
    automation_class.return_value.fill_item.assert_called_once()
    assert preparation.submitted is False
    assert preparation.prepared_count == 1
    assert preparation.city_code == "320100"


def test_service_keeps_unsubmitted_people_in_separate_tabs_and_continues_missing_person(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    page.context.new_page.return_value = Mock()
    logger = Mock()
    with patch(
        "ehrm.modules.person_information_collection.service.PersonInformationPage"
    ) as automation_class:
        automation_class.return_value.fill_item.side_effect = [
            EmployeeNotFoundError("人员不存在"), None
        ]
        result = PersonInformationService(settings, logger).prepare_with_page(
            page,
            [_item(), _item(identity_number="320101199002021235", source_index=3)],
        )
    page.context.new_page.assert_called_once()
    assert [row.success for row in result.results] == [False, True]
    assert result.submitted is False


def test_service_publishes_completed_rows_before_later_system_failure(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    logger = Mock()
    completed = []
    with patch(
        "ehrm.modules.person_information_collection.service.PersonInformationPage"
    ) as automation_class:
        automation_class.return_value.fill_item.side_effect = [
            None,
            WebsiteStructureChangedError("页面结构变化"),
        ]
        with pytest.raises(WebsiteStructureChangedError):
            PersonInformationService(
                settings, logger, result_callback=completed.append
            ).prepare_with_page(
                page,
                [_item(), _item(identity_number="320101199002021235", source_index=3)],
            )
    assert len(completed) == 1
    assert completed[0].item.source_index == 2


def test_cli_check_input_rejects_non_nanjing_configuration(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    settings = replace(
        settings,
        employment_termination=replace(
            settings.employment_termination,
            city_code="320200",
            city_name="无锡",
        ),
    )
    with (
        patch(
            "ehrm.entrypoints.person_information_collection_e2e_cli.application_runtime_root",
            return_value=tmp_path,
        ),
        patch(
            "ehrm.entrypoints.person_information_collection_e2e_cli.load_settings",
            return_value=settings,
        ),
    ):
        assert main(["--input", str(tmp_path / "not-needed.xlsx"), "--check-input"]) == 2


def test_cli_check_input_reuses_saved_default_account_without_browser(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    AuthenticationRepository(settings.auth_database_path).save_account(
        SystemType.JSHRSS,
        "test-unit-code",
        "test-password",
        secondary_account="test-mobile",
    )
    path = tmp_path / "采集测试.xlsx"
    workbook = Workbook()
    workbook.active.append(
        ["身份证号", "姓名", "手机号", "民族", "户籍性质", "省", "市", "区县", "街道", "社区村"]
    )
    workbook.active.append(
        ["320101199001011234", "测试人员", "13800138000", "汉族", "11", "江苏省", "南京市", "玄武区", "玄武门街道", "社区甲"]
    )
    workbook.save(path)
    workbook.close()
    with (
        patch(
            "ehrm.entrypoints.person_information_collection_e2e_cli.application_runtime_root",
            return_value=tmp_path,
        ),
        patch(
            "ehrm.entrypoints.person_information_collection_e2e_cli.load_settings",
            return_value=settings,
        ),
        patch("ehrm.entrypoints.person_information_collection_e2e_cli.BrowserManager") as browser,
    ):
        assert main(["--input", str(path), "--check-input"]) == 0
    browser.assert_not_called()
