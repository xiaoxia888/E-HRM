from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
import pytest

from ehrm.browser.smart_wait import SmartWait
from ehrm.core.error_catalog import ErrorCode
from ehrm.core.exceptions import (
    EmployeeNotFoundError,
    ExcelValidationError,
    QueryValidationError,
)
from ehrm.core.settings import load_settings
from ehrm.modules.employment_termination.excel_loader import (
    EmploymentTerminationExcelLoader,
)
from ehrm.modules.employment_termination.models import (
    EmploymentTerminationItem,
    EmploymentTerminationReason,
    normalize_termination_items,
)
from ehrm.modules.employment_termination.page import (
    EmploymentTerminationPage,
    _ResponseMonitor,
    _normalized_reason,
    _response_error_message,
)
from ehrm.modules.employment_termination.service import (
    EmploymentTerminationService,
)
from ehrm.modules.employment_termination.result_workbook import (
    EmploymentTerminationResultWorkbookWriter,
)


def _workbook(path: Path, rows: list[tuple[object, object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["身份证号", "退工原因"])
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def _immediate_pacer() -> Mock:
    pacer = Mock()
    pacer.perform.side_effect = lambda action: action()
    return pacer


def test_excel_adapter_converts_rows_to_domain_objects(tmp_path: Path) -> None:
    path = tmp_path / "退保测试.xlsx"
    _workbook(
        path,
        [
            ("320101199001011234", "6304"),
            ("320101199002021235", "单位解除合同"),
        ],
    )

    items = EmploymentTerminationExcelLoader().load(path)

    assert items == [
        EmploymentTerminationItem(
            "320101199001011234",
            EmploymentTerminationReason.MUTUAL_AGREEMENT,
            2,
        ),
        EmploymentTerminationItem(
            "320101199002021235",
            EmploymentTerminationReason.EMPLOYER_TERMINATED_CONTRACT,
            3,
        ),
    ]


def test_excel_adapter_requires_exact_two_business_columns(tmp_path: Path) -> None:
    path = tmp_path / "错误.xlsx"
    workbook = Workbook()
    workbook.active.append(["身份证", "原因"])
    workbook.save(path)
    workbook.close()

    with pytest.raises(ExcelValidationError, match="缺少必要列"):
        EmploymentTerminationExcelLoader().load(path)


def test_domain_object_rejects_duplicate_identity_numbers() -> None:
    items = [
        EmploymentTerminationItem("320101199001011234", "6301", 1),
        EmploymentTerminationItem("320101199001011234", "6304", 2),
    ]

    with pytest.raises(QueryValidationError, match="身份证号重复"):
        normalize_termination_items(items)


def test_reason_normalization_accepts_recorded_angular_value_and_plain_code() -> None:
    assert _normalized_reason("string:6304") == "6304"
    assert _normalized_reason(" 6304 ") == "6304"
    assert _normalized_reason("单位 解除劳动合同") == "单位解除劳动合同"


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("6301", EmploymentTerminationReason.CONTRACT_EXPIRED),
        ("string:6304", EmploymentTerminationReason.MUTUAL_AGREEMENT),
        ("单位解除合同", EmploymentTerminationReason.EMPLOYER_TERMINATED_CONTRACT),
        ("申请养老待遇核定（或达到法定退休年龄）", EmploymentTerminationReason.RETIREMENT_BENEFIT_APPLICATION),
        ("企业关闭或企业撤销、解散", EmploymentTerminationReason.ENTERPRISE_CLOSURE),
        ("申领病残津贴", EmploymentTerminationReason.DISABILITY_ALLOWANCE_APPLICATION),
    ],
)
def test_termination_reason_accepts_code_angular_value_or_name(
    value: str,
    expected: EmploymentTerminationReason,
) -> None:
    assert EmploymentTerminationReason.parse(value) is expected


def test_unknown_termination_reason_is_rejected_before_browser_is_open() -> None:
    with pytest.raises(QueryValidationError, match="不支持的退工原因"):
        EmploymentTerminationItem(
            "320101199001011234",
            "测试未知原因",
        ).normalized()


def test_network_business_errors_are_detected() -> None:
    assert _response_error_message({"appcode": "0", "msg": ""}) == ""
    assert (
        _response_error_message({"appcode": "1", "msg": "人员状态不允许退保"})
        == "人员状态不允许退保"
    )
    assert (
        _response_error_message(
            {"errors": [{"message": "Your session has expired."}]}
        )
        == "Your session has expired."
    )


def test_http_400_person_query_is_reported_as_employee_not_found() -> None:
    response = Mock()
    response.url = (
        "https://example.test/api/ehrss-si-enterprise-app/api/simis/persons/employees"
    )
    response.status = 400
    response.request.method = "GET"
    response.json.return_value = {"message": "未查询到人员信息"}
    item = EmploymentTerminationItem(
        "320101199001011234",
        "6304",
        source_index=2,
    )

    with pytest.raises(EmployeeNotFoundError, match="第 2 行.*未查询到人员信息"):
        EmploymentTerminationPage._validate_response(
            Mock(),
            response,
            item=item,
        )


def test_http_400_person_query_without_json_has_business_fallback() -> None:
    response = Mock()
    response.url = "https://example.test/api/persons/employees"
    response.status = 400
    response.request.method = "GET"
    response.json.side_effect = ValueError("not json")

    with pytest.raises(EmployeeNotFoundError, match="未查询到符合条件的参保人员"):
        EmploymentTerminationPage._validate_response(Mock(), response)


def test_response_monitor_ignores_unrelated_requests() -> None:
    unrelated = Mock()
    unrelated.url = "https://example.test/api/unrelated"
    unrelated.request.url = unrelated.url
    unrelated.request.post_data = "320101199001011234"
    expected = Mock()
    expected.url = "https://example.test/api/persons/employees?value=320101199001011234"
    expected.request.url = expected.url
    expected.request.post_data = None

    monitor = _ResponseMonitor.__new__(_ResponseMonitor)
    monitor.expected_path = "/api/persons/employees"
    monitor.responses = [unrelated, expected]
    monitor.waiter = SmartWait(lambda _milliseconds: None)

    response = monitor.wait_for_action_response(
        0,
        preferred_value="320101199001011234",
        timeout_ms=1000,
    )

    assert response is expected


def test_visible_region_prompt_selects_nanjing_before_navigation(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    option_locator = Mock()
    option = Mock()
    option.count.return_value = 1
    option.is_visible.return_value = True
    option_locator.count.return_value = 1
    option_locator.nth.return_value = option
    switch_locator = Mock()
    switch_locator.first = Mock()
    page.locator.side_effect = [option_locator, switch_locator]

    automation = EmploymentTerminationPage.__new__(EmploymentTerminationPage)
    automation.page = page
    automation.settings = settings
    automation.contract = settings.employment_termination
    automation.progress_callback = Mock()
    automation.cancel_check = None
    automation.pacer = _immediate_pacer()
    automation.waiter = SmartWait(page.wait_for_timeout)
    automation._read_region_state = Mock(
        side_effect=[
            {"areaCode": "320000", "cityname": "省本级"},
            {"areaCode": "320100", "cityname": "南京"},
        ]
    )

    automation._ensure_nanjing()

    page.locator.assert_any_call('[id="320100_1"]')
    option.click.assert_called_once_with()
    automation.progress_callback.assert_called_once_with(
        "退保：检测到办事地区选择窗口，正在选择南京"
    )


def test_absent_region_prompt_accepts_existing_nanjing_state(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    options = Mock()
    switch_locator = Mock()
    switch = Mock()
    switch_locator.first = switch
    page.locator.side_effect = [options, switch_locator]

    automation = EmploymentTerminationPage.__new__(EmploymentTerminationPage)
    automation.page = page
    automation.settings = settings
    automation.contract = settings.employment_termination
    automation.progress_callback = Mock()
    automation.cancel_check = None
    automation.pacer = _immediate_pacer()
    automation.waiter = SmartWait(page.wait_for_timeout)
    automation._read_region_state = Mock(
        return_value={"areaCode": "320100", "cityname": "南京"}
    )

    automation._ensure_nanjing()

    page.locator.assert_any_call('[id="320100_1"]')
    options.nth.assert_not_called()
    switch.click.assert_not_called()
    automation.progress_callback.assert_not_called()


def test_optional_entry_notice_falls_back_when_pointer_click_is_ignored(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    page.frames = []
    frame = Mock()
    search = Mock()
    search.count.return_value = 1
    search.is_visible.return_value = True
    search.is_enabled.return_value = True
    search_container = Mock()
    search_container.get_by_role.return_value.first = search
    dialogs = Mock()
    dialogs.filter.return_value = dialogs
    dialogs.count.return_value = 1
    dialog = Mock()
    dialog.count.return_value = 1
    dialog.is_visible.return_value = True
    dialog.inner_text.return_value = (
        "当前时间办理人员停保，发送税务缴费截止时间为2026年08月"
    )
    dialog_handle = Mock()
    dialog_handle.inner_text.return_value = dialog.inner_text.return_value
    dialog.element_handle.return_value = dialog_handle
    dialogs.nth.return_value = dialog
    button = Mock()
    button.count.return_value = 1
    button.is_visible.return_value = True
    buttons = Mock()
    buttons.count.return_value = 1
    buttons.nth.return_value = button
    button_handle = Mock()
    button.element_handle.return_value = button_handle
    dialog.locator.return_value = buttons
    frame.locator.side_effect = lambda selector: (
        search_container
        if selector == settings.employment_termination.person_search
        else dialogs
    )
    entry_confirms = Mock()
    entry_confirms.count.return_value = 0
    frame.get_by_role.return_value = entry_confirms

    automation = EmploymentTerminationPage.__new__(EmploymentTerminationPage)
    automation.page = page
    automation.settings = settings
    automation.contract = settings.employment_termination
    automation.progress_callback = Mock()
    automation.pacer = _immediate_pacer()
    current_time = [0.0]

    def advance(milliseconds: int) -> None:
        current_time[0] += milliseconds / 1000

    automation.waiter = SmartWait(advance, clock=lambda: current_time[0])
    dialog_closed = False

    def actionable(locator: object) -> bool:
        nonlocal dialog_closed
        if locator is search:
            return dialog_closed
        return locator is button

    # The real page has exhibited this exact behavior: Playwright reports that
    # the pointer click completed but Angular leaves the same modal in place.
    # Native HTMLElement.click() is the bounded fallback for that condition.
    def close_dialog(*_args: object, **_kwargs: object) -> None:
        nonlocal dialog_closed
        dialog_closed = True
        dialogs.count.configure_mock(return_value=0)
        dialog.is_visible.return_value = False

    dialog_handle.is_visible.side_effect = lambda: not dialog_closed
    button_handle.evaluate.side_effect = close_dialog

    automation._actionable = actionable

    automation._dismiss_entry_notice(frame)

    dialog.locator.assert_any_call(
        'button[data-ng-click="modalOptions.ok();"]'
    )
    button_handle.click.assert_called_once()
    assert button_handle.click.call_args.kwargs["timeout"] > 0
    assert "force" not in button_handle.click.call_args.kwargs
    button_handle.evaluate.assert_called_once_with(
        "element => { element.focus(); element.click(); }"
    )
    button.wait_for.assert_not_called()
    automation.progress_callback.assert_called_once_with(
        "退保：检测到停保时间提示，正在确认"
    )


def test_business_frame_uses_newest_visible_recorded_frame_chain(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = Mock()
    outer_frames = Mock()
    old_outer = Mock()
    new_outer = Mock()
    outer_frames.count.return_value = 2
    outer_frames.nth.side_effect = lambda index: [old_outer, new_outer][index]
    page.locator.return_value = outer_frames

    new_outer.count.return_value = 1
    new_outer.is_visible.return_value = True
    outer_handle = Mock()
    outer_frame = Mock()
    new_outer.element_handle.return_value = outer_handle
    outer_handle.content_frame.return_value = outer_frame

    business_frames = Mock()
    business_iframe = Mock()
    business_frames.count.return_value = 1
    business_frames.nth.return_value = business_iframe
    outer_frame.locator.return_value = business_frames
    business_iframe.count.return_value = 1
    business_iframe.is_visible.return_value = True
    business_handle = Mock()
    current_business_frame = Mock()
    business_iframe.element_handle.return_value = business_handle
    business_handle.content_frame.return_value = current_business_frame

    automation = EmploymentTerminationPage.__new__(EmploymentTerminationPage)
    automation.page = page
    automation.settings = settings
    automation.contract = settings.employment_termination
    automation.waiter = SmartWait(lambda _milliseconds: None)

    result = automation._wait_for_business_frame()

    assert result is current_business_frame
    page.locator.assert_called_once_with(
        'iframe[name^="layui-layer-iframe"]'
    )
    outer_frame.locator.assert_called_once_with("#busiIframe")
    old_outer.element_handle.assert_not_called()


def test_recorded_notice_and_entry_confirm_are_clicked_once_in_order(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    state = {"phase": "notice"}
    frame = Mock()

    search = Mock()
    search.count.return_value = 1
    search.is_visible.return_value = True
    search.is_enabled.return_value = True
    search_container = Mock()
    search_container.get_by_role.return_value.first = search

    dialogs = Mock()
    dialogs.count.side_effect = lambda: {
        "notice": 2,
        "entry": 1,
        "ready": 0,
    }[state["phase"]]
    notice_dialog = Mock()
    notice_dialog.count.return_value = 1
    notice_dialog.is_visible.side_effect = lambda: state["phase"] == "notice"
    notice_dialog.inner_text.return_value = (
        "当前时间办理人员停保，发送税务缴费截止时间为2026年08月"
    )
    notice_dialog_handle = Mock()
    notice_dialog_handle.is_visible.side_effect = (
        lambda: state["phase"] == "notice"
    )
    notice_dialog_handle.inner_text.return_value = notice_dialog.inner_text.return_value
    notice_dialog.element_handle.return_value = notice_dialog_handle
    entry_dialog = Mock()
    entry_dialog.count.return_value = 1
    entry_dialog.is_visible.side_effect = lambda: state["phase"] == "entry"
    entry_dialog.inner_text.return_value = "请确认进入退工停保登记"
    entry_dialog_handle = Mock()
    entry_dialog_handle.is_visible.side_effect = (
        lambda: state["phase"] == "entry"
    )
    entry_dialog_handle.inner_text.return_value = entry_dialog.inner_text.return_value
    entry_dialog.element_handle.return_value = entry_dialog_handle
    dialogs.nth.side_effect = lambda index: (
        notice_dialog
        if state["phase"] == "notice" and index == 1
        else entry_dialog
    )
    frame.locator.side_effect = lambda selector: (
        search_container
        if selector == settings.employment_termination.person_search
        else dialogs
    )
    notice_button = Mock()
    notice_button.count.return_value = 1
    notice_button.is_visible.return_value = True
    notice_buttons = Mock()
    notice_buttons.count.return_value = 1
    notice_buttons.nth.return_value = notice_button
    notice_dialog.locator.return_value = notice_buttons
    notice_button_handle = Mock()
    notice_button_handle.click.side_effect = (
        lambda **_kwargs: state.update(phase="entry")
    )
    notice_button.element_handle.return_value = notice_button_handle

    entry_button = Mock()
    entry_buttons = Mock()
    entry_buttons.count.return_value = 1
    entry_buttons.nth.return_value = entry_button
    entry_dialog.locator.return_value = entry_buttons
    entry_button_handle = Mock()
    entry_button_handle.click.side_effect = (
        lambda **_kwargs: state.update(phase="ready")
    )
    entry_button.element_handle.return_value = entry_button_handle

    automation = EmploymentTerminationPage.__new__(EmploymentTerminationPage)
    automation.page = Mock()
    automation.settings = settings
    automation.contract = settings.employment_termination
    automation.progress_callback = Mock()
    automation.pacer = _immediate_pacer()
    current_time = [0.0]

    def advance(milliseconds: int) -> None:
        current_time[0] += milliseconds / 1000

    automation.waiter = SmartWait(advance, clock=lambda: current_time[0])
    automation._actionable = lambda locator: (
        state["phase"] == "ready"
        if locator is search
        else (
            locator is notice_button and state["phase"] == "notice"
        )
        or (locator is entry_button and state["phase"] == "entry")
    )

    automation._dismiss_entry_notice(frame)

    notice_button_handle.click.assert_called_once()
    entry_button_handle.click.assert_called_once()
    assert automation.progress_callback.call_args_list == [
        (("退保：检测到停保时间提示，正在确认",), {}),
        (("退保：检测到业务入口确认，正在进入人员查询表单",), {}),
    ]


def test_visible_enabled_search_is_ready_after_all_dialogs_close(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    frame = Mock()
    search = Mock()
    search.count.return_value = 1
    search.is_visible.return_value = True
    search.is_enabled.return_value = True
    search_container = Mock()
    search_container.get_by_role.return_value.first = search
    dialogs = Mock()
    dialogs.count.return_value = 0
    frame.locator.side_effect = lambda selector: (
        search_container
        if selector == settings.employment_termination.person_search
        else dialogs
    )

    automation = EmploymentTerminationPage.__new__(EmploymentTerminationPage)
    automation.settings = settings
    automation.contract = settings.employment_termination
    automation.waiter = SmartWait(lambda _milliseconds: None)
    # A custom component may fail the centre-point hit test even when its
    # nested textbox is visible, enabled, and accepts fill().
    automation._actionable = Mock(return_value=False)

    automation._dismiss_entry_notice(frame)

    automation._actionable.assert_not_called()


def test_loading_selector_does_not_treat_layui_dialog_as_loading(
    tmp_path: Path,
) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)

    assert ".layui-layer-shade" not in settings.employment_termination.loading_indicator


def test_service_accepts_object_array_and_never_submits(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = object()
    frame = object()
    automation = Mock()
    automation.open.return_value = frame
    items = [EmploymentTerminationItem("320101199001011234", "6304")]

    with patch(
        "ehrm.modules.employment_termination.service.EmploymentTerminationPage",
        return_value=automation,
    ):
        result = EmploymentTerminationService(
            settings,
            Mock(),
        ).prepare_with_page(page, items)  # type: ignore[arg-type]

    automation.open.assert_called_once_with()
    automation.fill_item.assert_called_once_with(frame, items[0])
    assert result.prepared_count == 1
    assert result.failed_count == 0
    assert result.total_count == 1
    assert result.results[0].success is True
    assert result.city_code == "320100"
    assert result.city_name == "南京"
    assert result.submitted is False
    assert not hasattr(automation, "submit") or not automation.submit.called


def test_service_continues_after_employee_not_found(tmp_path: Path) -> None:
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    page = object()
    frame = object()
    automation = Mock()
    automation.open.return_value = frame
    automation.fill_item.side_effect = [
        EmployeeNotFoundError("未能找到该人员信息"),
        None,
    ]
    items = [
        EmploymentTerminationItem("320101199001011234", "6304", 2),
        EmploymentTerminationItem("320101199002021235", "6303", 3),
    ]

    with patch(
        "ehrm.modules.employment_termination.service.EmploymentTerminationPage",
        return_value=automation,
    ):
        result = EmploymentTerminationService(
            settings,
            Mock(),
        ).prepare_with_page(page, items)  # type: ignore[arg-type]

    assert automation.fill_item.call_count == 2
    automation.dismiss_query_feedback.assert_called_once_with(frame)
    assert result.total_count == 2
    assert result.prepared_count == 1
    assert result.failed_count == 1
    assert result.results[0].success is False
    assert result.results[0].code == ErrorCode.EMPLOYEE_NOT_FOUND
    assert result.results[0].message == "未能找到该人员信息"
    assert result.results[1].success is True


def test_result_workbook_contains_every_row_result(tmp_path: Path) -> None:
    source = tmp_path / "退保测试.xlsx"
    _workbook(
        source,
        [
            ("320101199001011234", "6304"),
            ("320101199002021235", "6303"),
        ],
    )
    settings = load_settings(Path("config/settings.toml"), data_root=tmp_path)
    automation = Mock()
    automation.open.return_value = object()
    automation.fill_item.side_effect = [
        EmployeeNotFoundError("未能找到该人员信息"),
        None,
    ]
    items = EmploymentTerminationExcelLoader().load(source)
    with patch(
        "ehrm.modules.employment_termination.service.EmploymentTerminationPage",
        return_value=automation,
    ):
        preparation = EmploymentTerminationService(
            settings,
            Mock(),
        ).prepare_with_page(object(), items)  # type: ignore[arg-type]

    destination = EmploymentTerminationResultWorkbookWriter().write(
        source,
        tmp_path / "results",
        preparation.results,
    )

    workbook = load_workbook(destination)
    try:
        sheet = workbook.active
        assert sheet.cell(1, 3).value == "处理结果"
        assert sheet.cell(1, 4).value == "失败原因"
        assert sheet.cell(2, 3).value == "未查询到人员"
        assert sheet.cell(2, 4).value == "未能找到该人员信息"
        assert sheet.cell(3, 3).value == "录入成功"
        assert sheet.cell(3, 4).value in (None, "")
    finally:
        workbook.close()
